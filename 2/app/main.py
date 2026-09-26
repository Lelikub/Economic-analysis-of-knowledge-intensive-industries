"""Единая точка входа расчётного pipeline практического занятия 2."""
from __future__ import annotations

import json
import logging
import subprocess
import sys
from dataclasses import asdict
from pathlib import Path

from docx import Document
from openpyxl import load_workbook

from src.harness import ValidationHarness
from src.loader import DataLoader
from src.metrics import ProjectMetricsCalculator
from src.reporting import ExcelReportExporter, WordReportGenerator


BASE = Path(__file__).resolve().parent


def _logger(path: Path) -> logging.Logger:
    logger = logging.getLogger("ea_practical_2")
    logger.setLevel(logging.INFO)
    for handler in list(logger.handlers):
        logger.removeHandler(handler)
        handler.close()
    handler = logging.FileHandler(path, encoding="utf-8", mode="w")
    handler.setFormatter(logging.Formatter("%(asctime)s | %(levelname)s | %(message)s"))
    logger.addHandler(handler)
    return logger


def _issues(metrics) -> list[dict[str, str]]:
    chain = metrics.chain
    first, last = chain["first_degraded"], chain["last_degraded"]
    issue_list = [
        {"code": "A01", "text": "Снижение КПД на 5% = умножение на 0,95", "reason": "формулировка задания не уточняет относительность", "effect": "сценарии цепочки"},
        {"code": "M01", "text": "Нет детальных m_net, f_scrap, p_i по материалам", "reason": "в BOM дана агрегированная сумма", "effect": "детальный скрап не рассчитан"},
        {"code": "M02", "text": "Нет h0/s0 и химического состава потоков", "reason": "не представлены в CSV/JSON/XLSX/лекции", "effect": "физическая и химическая эксергия отдельных потоков Н/Д"},
        {"code": "M03", "text": "Нет полного граничного набора потоков энергии", "reason": "не заданы электрическая мощность и отдельные теплопотери", "effect": "глобальный баланс и полный OPEX утилит Н/Д"},
        {"code": "M04", "text": "Нет c_F,1 цепочки", "reason": "sequential_chain_speco.csv содержит только эксергию и Z", "effect": "абсолютная стоимость цепочки и f_k Н/Д"},
        {"code": "M05", "text": "Нет коэффициента k для Ex_dest=k·ΔT", "reason": "лекция 3 задаёт пропорциональность без численного коэффициента", "effect": "оптимум теплообменника Н/Д"},
        {"code": "D01", "text": "В задании записано 1/0.85 ≈ 3,05", "reason": "математически 1/0,8⁵ = 3,05176", "effect": "в расчёте использована правильная степень"},
        {"code": "D02", "text": "Невязка котла-утилизатора 5 600 кВт", "reason": "газ отдаёт 47 600 кВт, пар воспринимает 42 000 кВт", "effect": "не классифицируется как доказанная потеря без дополнительного потока"},
        {"code": "D03", "text": "Последняя ступень цепочки округлена", "reason": "20,48 − 16,38 − 4,10 = 0,00 МВт; отношение 16,38/20,48 немного меньше 0,8", "effect": "допуск в проверке КПД 0,001"},
        {"code": "A02", "text": "Для чувствительности тарифа электричества взято 50–150% базы", "reason": "точный диапазон не задан", "effect": "только сценарный график"},
    ]
    # Вывод задания о максимальном вреде последней стадии проверяется на фактической модели.
    first_delta = first["capital_increment_usd_gj"] - chain["base"]["capital_increment_usd_gj"]
    last_delta = last["capital_increment_usd_gj"] - chain["base"]["capital_increment_usd_gj"]
    if abs(first_delta - last_delta) <= 1e-8:
        issue_list.append({"code": "D04", "text": "Положение ухудшенной ступени не меняет конечную стоимость в данной модели",
                           "reason": f"при фиксированных Ex_F,1 и Z_k оба сценария дают одинаковые A={first['amplification']:.6f} и ΔB={first_delta:.6f} USD/ГДж; произведение КПД коммутативно",
                           "effect": "утверждение задания о максимальном ущербе последней стадии не подтверждается исходной последовательной моделью"})
    elif first_delta > last_delta:
        issue_list.append({"code": "D04", "text": "Ущерб первой ступени выше ущерба последней по добавке B",
                           "reason": f"ΔB первого аппарата {first_delta:.6f} > ΔB пятого {last_delta:.6f} USD/ГДж",
                           "effect": "утверждение задания о максимальном ущербе последней стадии не подтверждается"})
    return issue_list


def _write_issues(path: Path, issues: list[dict[str, str]]) -> None:
    lines = ["# Замечания, допущения и расхождения", "", "Числовые результаты получены только расчётным кодом. Контрольные значения не подставлялись.", ""]
    for item in issues:
        lines.extend([f"## {item['code']} — {item['text']}", "", f"Основание: {item['reason']}.", f"Влияние: {item['effect']}.", ""])
    path.write_text("\n".join(lines), encoding="utf-8")


def _run_tests() -> str:
    command = [sys.executable, "-m", "pytest", "tests/test_calculations.py", "tests/test_loader.py", "-q", "-p", "no:cacheprovider"]
    result = subprocess.run(command, cwd=BASE, capture_output=True, text=True, encoding="utf-8", errors="replace")
    if result.returncode:
        raise RuntimeError(f"pytest завершился с ошибкой:\n{result.stdout}\n{result.stderr}")
    return result.stdout.strip().splitlines()[-1]


def _verify_artifacts(results: Path, logs: Path) -> None:
    names = [results / "metrics.json", results / "data_quality_report.json", results / "issues_and_assumptions.md",
             results / "EA_Practical_2_Report.xlsx", results / "EA_Practical_2_Report.docx",
             logs / "pipeline.log", logs / "harness_log.csv"]
    for file in names:
        if not file.is_file() or file.stat().st_size == 0:
            raise RuntimeError(f"Отсутствует или пуст артефакт: {file}")
    book = load_workbook(results / "EA_Practical_2_Report.xlsx", read_only=True, data_only=True)
    if len(book.sheetnames) < 4:
        raise RuntimeError("В Excel меньше четырёх листов")
    book.close()
    doc = Document(results / "EA_Practical_2_Report.docx")
    paragraphs = "\n".join(p.text for p in doc.paragraphs)
    for heading in ("1. Цель и задачи", "6. Результаты расчёта", "7. Harness Log", "8. Тестирование", "10.1. Допущения модели", "10.2. Data Quality Report", "12. Итоговые выводы"):
        if heading not in paragraphs:
            raise RuntimeError(f"Нет главы Word: {heading}")
    full_text = paragraphs + "\n" + "\n".join(c.text for table in doc.tables for row in table.rows for c in row.cells)
    if "[" in full_text or "]" in full_text:
        raise RuntimeError("В Word остались служебные заполнители")


def run_pipeline(data_dir: Path | None = None, task_dir: Path | None = None,
                 output_dir: Path | None = None) -> dict:
    """Выполняет чтение, расчёт, независимые проверки и все экспорты."""
    data_dir = Path(data_dir) if data_dir else BASE.parent / "data"
    task_dir = Path(task_dir) if task_dir else BASE.parent / "task"
    output_dir = Path(output_dir) if output_dir else BASE
    results, logs = output_dir / "results", output_dir / "logs"
    results.mkdir(parents=True, exist_ok=True)
    logs.mkdir(parents=True, exist_ok=True)
    logger = _logger(logs / "pipeline.log")
    logger.info("start | main | запуск | START | Единый pipeline")
    try:
        loader = DataLoader(data_dir)
        data = loader.load()
        for name, entry in loader.audit()["files"].items():
            logger.info("load | data_loader | файл | %s | %s: строк=%s, SHA256=%s, кодировка=%s",
                        entry["status"], name, entry["rows"], entry["sha256"], entry["encoding"])
            logger.info("audit | data_loader | структура | %s | %s: столбцы=%s, пропуски=%s, дубликаты=%s, диапазоны=%s",
                        entry["status"], name, entry["columns"], entry.get("missing_values", {}),
                        entry.get("duplicates", 0), entry.get("numeric_ranges", {}))
        metrics = ProjectMetricsCalculator(data).calculate()
        logger.info("calculate | metrics | OPEX/energy/exergy/SPECO/chain/compressors/heat | CALCULATED | единицы USD, MUSD, MW, GJ/h проверены")
        logger.info("task1 | opex | OPEX: Q·Σ(s_i·p_i)+ΣF_j | CALCULATED | аммиак база=%.4f USD/т; ХимМаш база=%.4f MUSD/шт",
                    metrics.opex["ammonia_base"]["unit_cost"], metrics.opex["machinery_base"]["unit_cost"])
        logger.info("task2 | materials | m_gross=m_net/(1-f_scrap); C_labor=Q·Σ(t·r) | PARTIALLY_CALCULATED | агрегированный BOM=%.4f MUSD/шт",
                    metrics.materials["aggregate_bom_per_unit_musd"])
        logger.info("task3 | energy | Σm_in·h+W+Q_in=Σm_out·h+Q_loss | PARTIALLY_CALCULATED | невязка КУ=%.1f кВт",
                    metrics.energy["whrb"]["residual_kw"])
        logger.info("task4 | exergy | Ex_F=Ex_P+Ex_dest+Ex_loss | PARTIALLY_CALCULATED | турбина η=%.4f; невязка=%.6f МВт",
                    metrics.exergy["cogen_efficiency"], metrics.exergy["cogen_balance_residual_mw"])
        logger.info("task5-6 | speco | SPECO: C_F+Z=C_P+C_loss | CALCULATED | F-rule=%.4f; P-rule=%.4f USD/ГДж",
                    metrics.speco["f_rule"]["electricity_cost_usd_gj"], metrics.speco["p_rule"]["product_cost_usd_gj"])
        logger.info("task7 | chain | c_P,N=A·c_F,1+B | PARTIALLY_CALCULATED | A=%.6f; B=%.6f USD/ГДж",
                    metrics.chain["base"]["amplification"], metrics.chain["base"]["capital_increment_usd_gj"])
        logger.info("task8 | compressors | Компрессоры: W=m1·w1+m2·w2 | VALIDATED | m1=%.4f кг/с; метод=%s; экономия=%.2f USD/год",
                    metrics.compressors["optimal_m1"], metrics.compressors["numeric_method"], metrics.compressors["annual_savings"])
        logger.info("task9 | heat_exchanger | TAC=CRF·CAPEX+τ_a·c_F·Ex_dest | %s | %s",
                    metrics.heat_exchanger["status"], metrics.heat_exchanger.get("missing", ""))
        issues = _issues(metrics)
        harness = ValidationHarness()
        rows = harness.validate(data, metrics)
        harness.write(logs / "harness_log.csv")
        for row in rows:
            if row["Status"] != "PASS":
                logger.warning("validate | harness | %s | %s | %s", row["Metric"], row["Status"], row["Comment"])
        (results / "metrics.json").write_text(json.dumps(asdict(metrics), ensure_ascii=False, indent=2, allow_nan=False), encoding="utf-8")
        (results / "data_quality_report.json").write_text(json.dumps(loader.audit(), ensure_ascii=False, indent=2), encoding="utf-8")
        _write_issues(results / "issues_and_assumptions.md", issues)
        test_result = _run_tests()
        logger.info("test | pytest | расчётные и загрузочные тесты | PASS | %s", test_result)
        ExcelReportExporter(metrics, data).export(results / "EA_Practical_2_Report.xlsx")
        logger.info("report | excel | экспорт | PASS | 4 листа")
        template = next(task_dir.glob("Шаблон_отчета_ЭА_краткий_v2.docx"), None)
        if template is None:
            raise FileNotFoundError("Отсутствует обязательный шаблон Word")
        WordReportGenerator(template, metrics, data, loader.audit(), rows, issues, test_result).generate(results / "EA_Practical_2_Report.docx")
        logger.info("report | word | экспорт | PASS | шаблон заполнен")
        for handler in logger.handlers:
            handler.flush()
        _verify_artifacts(results, logs)
        logger.info("finish | main | завершение | PASS | артефакты созданы и проверены")
        return asdict(metrics)
    except Exception:
        logger.exception("pipeline | main | исключение | FAIL | выполнение прервано")
        raise
    finally:
        for handler in list(logger.handlers):
            logger.removeHandler(handler)
            handler.close()


if __name__ == "__main__":
    run_pipeline()
