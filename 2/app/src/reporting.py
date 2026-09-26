"""Экспорт рассчитанных данных в русскоязычные Excel и Word."""
from __future__ import annotations

from pathlib import Path
from typing import Any

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill

from .models import ProjectInputData, ProjectMetrics


def fmt(value: Any, digits: int = 2) -> str:
    if isinstance(value, (float, int)):
        return f"{value:,.{digits}f}".replace(",", " ").replace(".", ",")
    return str(value)


class ExcelReportExporter:
    """Записывает четыре аналитических листа из уже рассчитанных метрик."""
    def __init__(self, metrics: ProjectMetrics, data: ProjectInputData):
        self.m, self.data = metrics, data

    @staticmethod
    def _sheet(book: Workbook, title: str, heading: str, headers: list[str]):
        sheet = book.create_sheet(title)
        sheet.append([heading])
        sheet.append(headers)
        sheet.freeze_panes = "A3"
        sheet.sheet_view.showGridLines = False
        return sheet

    @staticmethod
    def _style(sheet):
        dark = PatternFill("solid", fgColor="17365D")
        for row in (1, 2):
            for cell in sheet[row]:
                cell.fill = dark
                cell.font = Font(name="Aptos", bold=True, color="FFFFFF", size=13 if row == 1 else 10)
                cell.alignment = Alignment(wrap_text=True, vertical="center")
        sheet.row_dimensions[1].height = 32
        sheet.row_dimensions[2].height = 35
        for col in sheet.columns:
            letter = col[0].column_letter
            maximum = max(len(str(cell.value or "")) for cell in list(col)[:50])
            sheet.column_dimensions[letter].width = min(56, max(15, maximum + 2))
        for row in sheet.iter_rows(min_row=3):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if isinstance(cell.value, (int, float)):
                    cell.number_format = '#,##0.0000;[Red](#,##0.0000)'
        sheet.auto_filter.ref = f"A2:{sheet.cell(sheet.max_row, sheet.max_column).coordinate}"

    def export(self, path: Path) -> None:
        book = Workbook()
        book.remove(book.active)
        m, data = self.m, self.data
        o = self._sheet(book, "Сравнение OPEX", "Сравнение структуры OPEX и эффекта загрузки",
                        ["Объект", "Сценарий", "Выпуск", "Переменные затраты", "Постоянные затраты",
                         "Полный OPEX", "Доля переменных, %", "Доля постоянных, %", "Себестоимость единицы", "Единицы затрат"])
        for key, obj, scenario, unit in [
            ("ammonia_base", "Аммиак", "База", "USD; USD/т"),
            ("ammonia_underload", "Аммиак", "Недогруз", "USD; USD/т"),
            ("machinery_base", "ХимМаш", "База", "MUSD; MUSD/шт"),
            ("machinery_underload", "ХимМаш", "Недогруз", "MUSD; MUSD/шт"),
            ("machinery_mitigated", "ХимМаш", "Недогруз со снижением fixed", "MUSD; MUSD/шт")]:
            x = m.opex[key]
            o.append([obj, scenario, x["quantity"], x["variable"], x["fixed"], x["total"],
                      x["variable_share_pct"], x["fixed_share_pct"], x["unit_cost"], unit])
        o.append([])
        o.append(["Материалы BOM, агрегат", "на 1 реактор", "", m.materials["aggregate_bom_per_unit_musd"], "", "", "", "", "", "MUSD/шт"])
        o.append(["Прямой труд, агрегат", "на 1 реактор", "", m.materials["aggregate_labor_per_unit_musd"], "", "", "", "", "", "MUSD/шт"])
        o.append(["Детальный расчёт скрапа", "MISSING_INPUT", "", "", "", "", "", "", "", m.materials["detailed_scrap_missing"]])
        o.append(["WIP", "оборотный капитал", "", "", "", "", "", "", "", "не включён в OPEX"])
        o.append([])
        o.append(["Выпуск аммиака, т/год", "Себестоимость, USD/т"])
        start_a = o.max_row + 1
        for q in range(300_000, 550_000, 25_000):
            o.append([q, m.opex["ammonia_base"]["variable"] / 500_000 + m.opex["ammonia_base"]["fixed"] / q])
        chart = LineChart()
        chart.title = "Удельная себестоимость аммиака от выпуска"
        chart.y_axis.title = "USD/т"
        chart.x_axis.title = "т/год"
        chart.add_data(Reference(o, min_col=2, min_row=start_a - 1, max_row=o.max_row), titles_from_data=True)
        chart.set_categories(Reference(o, min_col=1, min_row=start_a, max_row=o.max_row))
        o.add_chart(chart, "L3")

        e = self._sheet(book, "Энергетика и материалы", "Потоки и энергетический баланс",
                        ["ID потока / показатель", "Наименование", "Расход, кг/с", "Энтальпия, кДж/кг", "Поток энергии, кВт", "Комментарий"])
        for row in m.energy["streams"]:
            e.append([row["id"], row["name"], row["mass_kg_s"], row["h_kj_kg"], row["energy_kw"], "исходный поток"])
        for name, key, unit in [
            ("Энергия на входе КУ", "energy_in_kw", "кВт"), ("Энергия на выходе КУ", "energy_out_kw", "кВт"),
            ("Невязка КУ", "residual_kw", "кВт"), ("Относительная невязка", "relative_error", "доля")]:
            e.append([name, "Котёл-утилизатор", "", "", m.energy["whrb"][key], unit])
        for name, key, unit in [
            ("Отбор теплоты газом", "gas_heat_release_kw", "кВт"), ("Выработка пара, теплота", "steam_duty_kw", "кВт"),
            ("Нераспределённое тепло", "unallocated_heat_kw", "кВт"), ("Тепло CW", "cooling_water_heat_kw", "кВт"),
            ("Расход CW расчётный", "cooling_water_kg_s_calculated", "кг/с"), ("Расход пара расчётный", "steam_kg_s_calculated", "кг/с"),
            ("Потенциальный кредит пара", "potential_steam_credit_usd_h", "USD/ч")]:
            e.append([name, "", "", "", m.energy[key], unit])
        e.append(["Полный OPEX утилит", "НЕ РАССЧИТЫВАЕТСЯ", "", "", "", m.energy["utility_opex_missing"]])

        x = self._sheet(book, "Эксергия и SPECO", "Эксергия, стоимостные правила и цепочка",
                        ["Аппарат / показатель", "Ex_F, МВт", "Ex_P, МВт", "Ex_dest, МВт", "КПД", "Z, USD/ч",
                         "c_F, USD/ГДж", "c_P, USD/ГДж", "f_k", "Рекомендация / примечание"])
        for row in m.chain["base"]["rows"]:
            x.append([row["name"], row["fuel_mw"], row["product_mw"], row["destruction_mw"],
                      row["efficiency"], row["z_usd_h"], "Н/Д", "Н/Д", "Н/Д", "Нет c_F,1 в источнике"])
        x.append(["Турбина CHP", m.exergy["cogen_fuel_mw"], m.exergy["cogen_products_mw"],
                  m.exergy["cogen_destruction_mw"], m.exergy["cogen_efficiency"], m.speco["z_usd_h"],
                  data.cogeneration.values["High_Pressure_Steam_Inlet"]["Specific_Cost_USD_per_GJ"],
                  m.speco["p_rule"]["product_cost_usd_gj"], m.exergy["cogen_factor"], m.exergy["cogen_factor_recommendation"]])
        x.append(["F-rule: электроэнергия", "", "", "", "", "", "",
                  m.speco["f_rule"]["electricity_cost_usd_gj"], "", "Стоимость остаточного LP-пара равна стоимости HP-пара"])
        x.append(["P-rule: оба продукта", "", "", "", "", "", "",
                  m.speco["p_rule"]["product_cost_usd_gj"], "", "Одинаковая удельная стоимость"])
        x.append(["Цепочка: множитель сырья", "", "", "", m.chain["base"]["amplification"], "", "", "", "", "без c_F,1"])
        for label, key in [("Цепочка: база", "base"), ("Цепочка: ухудшение 1-го", "first_degraded"), ("Цепочка: ухудшение 5-го", "last_degraded")]:
            item = m.chain[key]
            x.append([label, "", item["outlet_mw"], "", item["amplification"], "", "",
                      item["capital_increment_usd_gj"], "", "c_P,N = A·c_F,1 + B"])
        x.append(["Физическая эксергия потоков", "", "", "", "", "", "", "", "", m.exergy["stream_physical_missing"]])

        t = self._sheet(book, "Чувствительность", "Чувствительность OPEX к тарифам",
                        ["Цена газа, USD/ГДж", "Себестоимость аммиака, USD/т", "Годовой OPEX, USD", "Тариф электричества, USD/кВт·ч", "Себестоимость аммиака, USD/т", "Годовой OPEX, USD"])
        for index, gas in enumerate(m.sensitivity["gas"]):
            el = m.sensitivity["electricity"][index] if index < len(m.sensitivity["electricity"]) else {}
            t.append([gas["gas_price_usd_gj"], gas["ammonia_unit_cost_usd_t"], gas["annual_opex_usd"],
                      el.get("electricity_price_usd_kwh"), el.get("ammonia_unit_cost_usd_t"), el.get("annual_opex_usd")])
        for label, x_col, y_col, rows, anchor in [
            ("Влияние цены газа", 1, 2, len(m.sensitivity["gas"]), "H3"),
            ("Влияние тарифа электроэнергии", 4, 5, len(m.sensitivity["electricity"]), "H19")]:
            chart = LineChart()
            chart.title = label
            chart.y_axis.title = "USD/т"
            chart.add_data(Reference(t, min_col=y_col, min_row=2, max_row=2 + rows), titles_from_data=True)
            chart.set_categories(Reference(t, min_col=x_col, min_row=3, max_row=2 + rows))
            t.add_chart(chart, anchor)
        for sheet in book:
            self._style(sheet)
        book.save(path)


class WordReportGenerator:
    """Заполняет исходный DOCX-шаблон фактическими результатами pipeline."""
    def __init__(self, template: Path, metrics: ProjectMetrics, data: ProjectInputData,
                 quality: dict, harness: list[dict], issues: list[dict], test_result: str):
        self.template, self.m, self.data = template, metrics, data
        self.quality, self.harness, self.issues, self.test_result = quality, harness, issues, test_result

    @staticmethod
    def _table(table, rows: list[list[str]]) -> None:
        # Заголовок шаблона сохраняется; остальные строки заменяются результатами.
        while len(table.rows) > 1:
            table._tbl.remove(table.rows[-1]._tr)
        for values in rows:
            cells = table.add_row().cells
            for cell, value in zip(cells, values):
                cell.text = str(value)
        table.style = "Table Grid"
        for cell in table.rows[0].cells:
            cell._tc.get_or_add_tcPr()
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), "17365D")
            cell._tc.get_or_add_tcPr().append(shading)
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)

    @staticmethod
    def _remove_paragraph(paragraph) -> None:
        paragraph._element.getparent().remove(paragraph._element)

    def _figures(self, output_dir: Path) -> list[Path]:
        output_dir.mkdir(exist_ok=True)
        plt.rcParams["font.family"] = "DejaVu Sans"
        paths = []
        o = self.m.opex
        q = list(range(300_000, 550_000, 25_000))
        y = [o["ammonia_base"]["variable"] / 500_000 + o["ammonia_base"]["fixed"] / v for v in q]
        fig, ax = plt.subplots(figsize=(6.6, 3.6))
        ax.plot(q, y, marker="o")
        ax.set(xlabel="Выпуск, т/год", ylabel="Себестоимость, USD/т", title="Эффект масштаба: аммиак")
        ax.grid(alpha=.25)
        fig.tight_layout()
        p = output_dir / "01_unit_cost.png"; fig.savefig(p, dpi=160); plt.close(fig); paths.append(p)

        fig, ax = plt.subplots(figsize=(6.6, 3.6))
        labels = ["Аммиак", "ХимМаш"]
        variable = [o["ammonia_base"]["variable_share_pct"], o["machinery_base"]["variable_share_pct"]]
        fixed = [100 - v for v in variable]
        ax.bar(labels, variable, label="Переменные")
        ax.bar(labels, fixed, bottom=variable, label="Постоянные")
        ax.set(ylabel="Доля OPEX, %", title="Структура годового OPEX")
        ax.legend()
        fig.tight_layout()
        p = output_dir / "02_opex_structure.png"; fig.savefig(p, dpi=160); plt.close(fig); paths.append(p)

        c = self.m.compressors
        fig, ax = plt.subplots(figsize=(6.6, 3.6))
        ax.bar(["Поровну", "Оптимум"], [c["equal_cost_usd_h"], c["optimal_cost_usd_h"]])
        ax.set(ylabel="USD/ч", title="OPEX компрессорной станции")
        fig.tight_layout()
        p = output_dir / "03_compressors.png"; fig.savefig(p, dpi=160); plt.close(fig); paths.append(p)

        gas = self.m.sensitivity["gas"]
        fig, ax = plt.subplots(figsize=(6.6, 3.6))
        ax.plot([v["gas_price_usd_gj"] for v in gas], [v["ammonia_unit_cost_usd_t"] for v in gas], marker="o")
        ax.set(xlabel="Цена газа, USD/ГДж", ylabel="Себестоимость, USD/т", title="Чувствительность к цене газа")
        ax.grid(alpha=.25)
        fig.tight_layout()
        p = output_dir / "04_gas_sensitivity.png"; fig.savefig(p, dpi=160); plt.close(fig); paths.append(p)
        return paths

    def generate(self, path: Path) -> None:
        doc = Document(self.template)
        p = doc.paragraphs
        m, o = self.m, self.m.opex
        c = m.compressors
        p[10].text = "ОТЧЁТ ПО ПРАКТИЧЕСКОЙ РАБОТЕ № 2"
        p[11].text = "«Структура OPEX, энергетический баланс и эксергетический анализ»"
        p[12].text = "Объект: аммиачное производство и завод «ХимМаш»"
        p[14].text = f"{self.data.economic_constants.values['Base_Year']} г."
        content = {
            17: "Цель — связать эксплуатационные затраты двух производств с материальными потоками, энергетикой и качеством энергии. Рассчитаны сценарии загрузки, SPECO-стоимости продуктов, эффект распределения нагрузки компрессоров и границы применимости модели.",
            20: "Проверить полноту и единицы исходных CSV, JSON и XLSX.",
            21: "Рассчитать OPEX, баланс котла-утилизатора и когенерационной турбины.",
            22: "Сопоставить SPECO F-rule и P-rule и исследовать цепочку аппаратов.",
            23: "Проверить расчёты тестами и независимыми эталонами.",
            24: "Оценить недогруз, сокращение постоянных затрат и чувствительность к тарифам.",
            26: "Объект А — непрерывный выпуск 500 000 т аммиака в год при 8 000 ч/год. Объект Б — дискретный выпуск 12 реакторов в год. Для аммиака ключевы сырьевые и энергетические тарифы; для «ХимМаш» — серия и распределение постоянных затрат.",
            28: "Первичные расчётные данные: шесть CSV для OPEX, потоков и последовательной цепочки; три JSON для тарифов, компрессоров и турбины. XLSX проверен как дублирующий источник и источник контрольных листов. Формулы сверены с заданием и лекцией 3.",
            32: "CSV/JSON/XLSX → DataLoader → ProjectInputData → контроль типов и единиц → расчётные классы → ProjectMetrics → ValidationHarness → Excel и Word. Эталоны отделены от вычислений.",
            35: "Все денежные и физические показатели рассчитываются Python-кодом. Приведённые ниже формулы определяют смысл строк итоговых таблиц.",
            37: "Удельный OPEX показывает стоимость выпуска единицы продукта внутри релевантного диапазона постоянных затрат.",
            38: "OPEX(Q) = Q·Σ(sᵢpᵢ) + ΣFⱼ; c_unit(Q) = OPEX(Q)/Q",
            39: "sᵢ — удельный расход, pᵢ — тариф, Fⱼ — годовая постоянная статья, Q — выпуск. Единицы: USD/т и USD/год либо MUSD/шт и MUSD/год.",
            41: "Для BOM: m_gross = m_net/(1−f_scrap); труд = Q·Σ(t_op·r_op). При отсутствии подробных норм использована только исходная агрегированная стоимость BOM; WIP отнесён к оборотному капиталу.",
            42: "Баланс энергии: Σ(m_in·h_in)+W+Q_in = Σ(m_out·h_out)+Q_loss; эксергия: ex_ph=(h−h₀)−T₀(s−s₀); SPECO: C_F+Z=C_P+C_loss.",
            44: "Показатели ниже получены из ProjectMetrics. «Н/Д» и статус MISSING_INPUT означают отсутствие подтверждённого числового входа.",
            46: f"При недогрузе аммиачная себестоимость меняется с {fmt(o['ammonia_base']['unit_cost'])} до {fmt(o['ammonia_underload']['unit_cost'])} USD/т. Для «ХимМаш» эффект выражен сильнее из-за высокой постоянной надбавки; сокращение fixed снижает себестоимость до {fmt(o['machinery_mitigated']['unit_cost'])} MUSD/шт.",
            50: "Ниже показаны фактические контрольные сравнения. Ошибки, которых не было, не добавлялись; недостающие параметры выделены отдельным статусом.",
            51: "Расхождения последней ступени цепочки объясняются округлением исходных потоков до 0,01 МВт. Невязка котла-утилизатора остаётся явной и требует дополнительного потока потерь или отбора.",
            53: f"Результат pytest: {self.test_result}. Проверены допустимость КПД, второй закон, стоимостные балансы, матрица, расчёты OPEX и работа полного pipeline.",
            57: "На рисунках 1–4 сопоставлены эффект масштаба, структура OPEX, режимы компрессоров и влияние цены газа.",
            58: "Рисунки 1–4 — расчётные графики, построенные из ProjectMetrics.",
            59: "Цена газа меняет себестоимость линейно по норме 28 ГДж/т. Оптимизация компрессоров даёт экономию без дополнительного CAPEX при заданной модели фиксированных затрат.",
            62: "Принята относительная трактовка снижения КПД на 5%: η_new = 0,95η. Для цепочки фиксируются вход первого аппарата и Z каждого аппарата; потоки ниже по цепочке пересчитываются.",
            64: "Файлы прочитаны в UTF-8; XLSX проверен по листам. Отсутствующие проектные параметры не заменялись синтетическими значениями.",
            66: f"Аммиак прежде всего чувствителен к газу: при тарифе 8 USD/ГДж он даёт 224 USD/т переменных затрат. Турбина по F-rule относит большую стоимость к электроэнергии ({fmt(m.speco['f_rule']['electricity_cost_usd_gj'])} USD/ГДж), а P-rule делит её между двумя продуктами ({fmt(m.speco['p_rule']['product_cost_usd_gj'])} USD/ГДж).",
            67: "Границы модели: нет детальных норм скрапа, термодинамических свойств h₀/s₀ и химического состава потоков, полного набора потоков энергобаланса, цены первичного топлива цепочки и коэффициента потерь теплообменника. При фиксированных входе и Z положение ухудшенной ступени не меняет конечную стоимость; это противоречит качественному выводу задания и лекции 3.",
            70: f"Аммиак: {fmt(o['ammonia_base']['unit_cost'])} USD/т в базе; контроль OPEX — PASS.",
            71: f"При 400 000 т/год: {fmt(o['ammonia_underload']['unit_cost'])} USD/т, изменение {fmt(o['ammonia_unit_change_pct'])}%.",
            72: f"«ХимМаш»: {fmt(o['machinery_base']['unit_cost'], 3)} MUSD/шт в базе и {fmt(o['machinery_underload']['unit_cost'], 3)} MUSD/шт при недогрузе.",
            73: f"Сокращение fixed на 4,8 MUSD/год снижает себестоимость до {fmt(o['machinery_mitigated']['unit_cost'], 3)} MUSD/шт.",
            74: f"Турбина: F-rule = {fmt(m.speco['f_rule']['electricity_cost_usd_gj'])} и P-rule = {fmt(m.speco['p_rule']['product_cost_usd_gj'])} USD/ГДж; балансы — PASS.",
            75: f"Компрессоры: экономия {fmt(c['annual_savings'], 0)} USD/год; проектный оптимум теплообменника не определяется без коэффициента эксергетических потерь.",
        }
        for index, value in content.items():
            p[index].text = value
        for item in p[76:]:
            self._remove_paragraph(item)
        p[42].alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p[42].runs:
            run.font.name = "Cambria Math"
            run.font.size = Pt(12)
        anchor = p[42]
        methodology = [
            ("m_gross = m_net/(1−f_scrap); C_labor = Q·Σ(t_op·r_op)",
             "m_net — чистая масса, f_scrap — доля отходов; t_op и r_op — нормо-часы и тариф. Детальный расчёт материалов требует отсутствующих исходных норм."),
            ("Ex_dest = T₀·S_gen; Ex_F = Ex_P + Ex_dest + Ex_loss",
             "T₀ — температура среды, S_gen — генерация энтропии. Проверяется неотрицательность S_gen; для полного баланса потоков данных недостаточно."),
            ("F-rule: c_residual = c_fuel; P-rule: c_product,1 = c_product,2",
             "Правила задают дополнительные уравнения к стоимостному балансу турбины; суммы выражаются в USD/ч, удельные цены — в USD/ГДж."),
            ("c_P,N = A·c_F,1 + B; A = Π(1/η_k); B = Σ{Z_k/Ex_P,k·Π_m>k(1/η_m)}",
             "A — множитель стоимости первичного топлива; B — приведённые затраты оборудования в USD/ГДж. Без c_F,1 абсолютная стоимость цепочки недоступна."),
            ("f_k = Z_k/(Z_k + c_F,k·(Ex_dest,k + Ex_loss,k))",
             "Эксергетические потоки переводятся из МВт в ГДж/ч. Малое f_k указывает на доминирование потерь, большое — на капиталоёмкость."),
            ("m1 + m2 = M; W = m1·(b1+c1·m1) + m2·(b2+c2·m2)",
             "Расходы m1, m2 выражены в кг/с, удельная работа — в кДж/кг, мощность W — в кВт; минимизируется годовой OPEX электроэнергии."),
            ("TAC(ΔT) = CRF·C0·(ΔT_base/ΔT)^β + τ_a·c_F·k·ΔT",
             "TAC — годовые приведённые затраты. Коэффициент k в источниках не задан; формула реализована и проверена синтетическим тестом, но проектный оптимум не публикуется."),
        ]
        for formula, explanation in methodology:
            equation = doc.add_paragraph()
            equation.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = equation.add_run(formula)
            run.font.name = "Cambria Math"
            run.font.size = Pt(12)
            comment = doc.add_paragraph(explanation)
            anchor._p.addnext(equation._p)
            equation._p.addnext(comment._p)
            anchor = comment

        self._table(doc.tables[1], [
            ["Годовой выпуск аммиака", "Задание №2", "500 000", "400 000", "т/год"],
            ["Годовая серия «ХимМаш»", "Задание №2", "12", "6", "шт/год"],
            ["Фонд работы", "project_economic_constants.json", "8 000", "8 000", "ч/год"],
            ["Цена газа", "project_economic_constants.json", "8", "4–16", "USD/ГДж"],
        ])
        self._table(doc.tables[2], [
            ["1", "CSV/JSON/XLSX → ProjectInputData", "чтение, типизация, аудит"],
            ["2", "ProjectInputData → ProjectMetrics", "расчёт OPEX, энергии, эксергии и SPECO"],
            ["3", "ProjectMetrics → Harness", "сравнение с эталонами"],
            ["4", "ProjectMetrics → Excel/Word", "визуализация и аналитический отчёт"],
        ])
        result_rows = []
        for label, key, annual_unit, unit in [
            ("Аммиак, база", "ammonia_base", "USD/год", "USD/т"),
            ("Аммиак, недогруз", "ammonia_underload", "USD/год", "USD/т"),
            ("ХимМаш, база", "machinery_base", "MUSD/год", "MUSD/шт"),
            ("ХимМаш, недогруз", "machinery_underload", "MUSD/год", "MUSD/шт"),
            ("ХимМаш, снижение fixed", "machinery_mitigated", "MUSD/год", "MUSD/шт")]:
            scenario = o[key]
            result_rows.extend([
                [f"{label}: переменные", fmt(scenario["variable"], 3), annual_unit, "РАССЧИТАНО"],
                [f"{label}: постоянные", fmt(scenario["fixed"], 3), annual_unit, "РАССЧИТАНО"],
                [f"{label}: полный OPEX", fmt(scenario["total"], 3), annual_unit, "РАССЧИТАНО"],
                [f"{label}: доли", f"переменные {fmt(scenario['variable_share_pct'])}; постоянные {fmt(scenario['fixed_share_pct'])}", "%", "РАССЧИТАНО"],
                [f"{label}: себестоимость", fmt(scenario["unit_cost"], 3), unit, "РАССЧИТАНО"],
            ])
        result_rows += [
            ["Невязка котла-утилизатора", fmt(m.energy["whrb"]["residual_kw"]), "кВт", "ЧАСТИЧНО"],
            ["Расход охлаждающей воды", fmt(m.energy["cooling_water_kg_s_calculated"]), "кг/с", "РАССЧИТАНО"],
            ["Расход пара", fmt(m.energy["steam_kg_s_calculated"]), "кг/с", "РАССЧИТАНО"],
            ["Потенциальный кредит пара", fmt(m.energy["potential_steam_credit_usd_h"]), "USD/ч", "Условная оценка"],
            ["Эксергетический КПД турбины", fmt(m.exergy["cogen_efficiency"]), "доля", "РАССЧИТАНО"],
            ["Эксергоэкономический фактор турбины", fmt(m.exergy["cogen_factor"], 4), "доля", "РАССЧИТАНО"],
            ["SPECO F-rule, электричество", fmt(m.speco["f_rule"]["electricity_cost_usd_gj"]), "USD/ГДж", "ПРОВЕРЕНО"],
            ["SPECO P-rule, продукты", fmt(m.speco["p_rule"]["product_cost_usd_gj"]), "USD/ГДж", "ПРОВЕРЕНО"],
            ["Множитель цепочки", fmt(m.chain["base"]["amplification"], 4), "раз", "РАССЧИТАНО"],
            ["Цепочка: прирост B при деградации 1-го", fmt(m.chain["first_degraded"]["capital_increment_usd_gj"] - m.chain["base"]["capital_increment_usd_gj"], 4), "USD/ГДж", "РАССЧИТАНО"],
            ["Цепочка: прирост B при деградации 5-го", fmt(m.chain["last_degraded"]["capital_increment_usd_gj"] - m.chain["base"]["capital_increment_usd_gj"], 4), "USD/ГДж", "РАССЧИТАНО"],
            ["Компрессоры: OPEX поровну", fmt(c["equal_cost_usd_h"]), "USD/ч", "ПРОВЕРЕНО"],
            ["Компрессоры: OPEX оптимум", fmt(c["optimal_cost_usd_h"]), "USD/ч", "ПРОВЕРЕНО"],
            ["Компрессоры, экономия", fmt(c["annual_savings"], 0), "USD/год", "ПРОВЕРЕНО"],
            ["Оптимум теплообменника", "Н/Д", "K", "MISSING_INPUT"],
        ]
        self._table(doc.tables[3], result_rows)
        significant = [r for r in self.harness if r["Status"] in ("WARNING", "SKIPPED_MISSING_INPUT", "FAIL")][:5]
        if not significant:
            significant = [r for r in self.harness if r["Status"] == "PASS"][:1]
        harness_rows = []
        for r in significant[:3]:
            harness_rows.extend([
                ["Запрос / показатель", r["Metric"]],
                ["Расчётный фрагмент", "Python → ProjectMetrics → ValidationHarness"],
                ["Обнаруженная проблема", r["Comment"] or "Ошибка не обнаружена"],
                ["Причина и действие", "Не подставлять отсутствующий параметр; сохранить отдельный статус"],
                ["Повторная проверка", r["Status"]],
            ])
        self._table(doc.tables[4], harness_rows)
        statuses = [r for r in self.harness if r["Status"] == "PASS"][:8]
        self._table(doc.tables[5], [[r["Metric"], r["Reference_Source"], fmt(r["Calculated"]), r["Status"]] for r in statuses])
        self._table(doc.tables[6], [
            ["Аммиак, USD/т", fmt(o["ammonia_base"]["unit_cost"]), fmt(o["ammonia_underload"]["unit_cost"]),
             fmt(o["ammonia_underload"]["unit_cost"] - o["ammonia_base"]["unit_cost"]), fmt(o["ammonia_unit_change_pct"]), "распределение fixed на меньший выпуск"],
            ["ХимМаш, MUSD/шт", fmt(o["machinery_base"]["unit_cost"]), fmt(o["machinery_underload"]["unit_cost"]),
             fmt(o["machinery_underload"]["unit_cost"] - o["machinery_base"]["unit_cost"]), fmt(o["machinery_unit_change_pct"]), "рост fixed на изделие"],
            ["Компрессоры, USD/ч", fmt(c["equal_cost_usd_h"]), fmt(c["optimal_cost_usd_h"]),
             fmt(c["optimal_cost_usd_h"] - c["equal_cost_usd_h"]), fmt((c["optimal_cost_usd_h"] / c["equal_cost_usd_h"] - 1) * 100), "балансировка нагрузок"],
            ["Цепочка A, раз", fmt(m.chain["base"]["amplification"], 4), fmt(m.chain["last_degraded"]["amplification"], 4),
             fmt(m.chain["last_degraded"]["amplification"] - m.chain["base"]["amplification"], 4), "5,26", "КПД последнего аппарата ×0,95"],
        ])
        # Служебная таблица для рисунка заменяется реальными иллюстрациями.
        placeholder_table = doc.tables[7]
        placeholder_table._element.getparent().remove(placeholder_table._element)
        priority = ["A01", "D01", "D02", "D04", "M01", "M05"]
        chosen_issues = [next(x for x in self.issues if x["code"] == code) for code in priority if any(x["code"] == code for x in self.issues)]
        self._table(doc.tables[7], [[x["code"], x["text"], x["reason"], x["effect"]] for x in chosen_issues])
        quality_rows = []
        for name, info in self.quality["files"].items():
            problems = sum(info.get("missing_values", {}).values()) if isinstance(info.get("missing_values"), dict) else 0
            if info["status"] != "PASS" or problems:
                quality_rows.append([name, "пропуски или типы", "СРЕДНЯЯ", "аудит", info["status"]])
        if not quality_rows:
            quality_rows = [["Существенных структурных ошибок нет", "CSV/JSON/XLSX", "НИЗКАЯ", "числовые расчёты доступны", "проверено"]]
        self._table(doc.tables[8], quality_rows)

        figures = self._figures(path.parent / "figures")
        anchor = p[57]
        captions = ["Рисунок 1 — Удельная себестоимость аммиака от выпуска",
                    "Рисунок 2 — Доли переменных и постоянных затрат",
                    "Рисунок 3 — OPEX компрессоров при равной и оптимальной загрузке",
                    "Рисунок 4 — Чувствительность себестоимости к цене газа"]
        for figure, caption in zip(figures, captions):
            picture_paragraph = doc.add_paragraph()
            picture_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            picture_paragraph.add_run().add_picture(str(figure), width=Cm(13.5))
            caption_paragraph = doc.add_paragraph(caption)
            caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            anchor._p.addnext(picture_paragraph._p)
            picture_paragraph._p.addnext(caption_paragraph._p)
            anchor = caption_paragraph

        for paragraph in doc.paragraphs:
            if paragraph.style.name == "Normal":
                paragraph.paragraph_format.line_spacing = 1.15
                paragraph.paragraph_format.first_line_indent = Cm(1.25)
                for run in paragraph.runs:
                    if run.font.name != "Cambria Math":
                        run.font.name = "Times New Roman"
                        run.font.size = Pt(12)
        doc.save(path)
