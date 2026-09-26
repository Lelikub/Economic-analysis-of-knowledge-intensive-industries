"""Проверка одного запуска от реальных входов до проверяемых артефактов."""
import csv
import json
from pathlib import Path

from docx import Document
from openpyxl import load_workbook

from main import run_pipeline


DATA = Path(__file__).resolve().parents[2] / "data"
TASK = Path(__file__).resolve().parents[2] / "task"


def test_full_pipeline_builds_russian_reports(tmp_path):
    run_pipeline(data_dir=DATA, task_dir=TASK, output_dir=tmp_path)
    results, logs = tmp_path / "results", tmp_path / "logs"
    metrics = json.loads((results / "metrics.json").read_text(encoding="utf-8"))
    assert metrics["opex"]["ammonia_base"]["unit_cost"] == 347.6
    assert metrics["exergy"]["cogen_factor"] > 0
    assert metrics["heat_exchanger"]["status"] == "SKIPPED_MISSING_INPUT"
    assert "положение ухудшенной ступени" in (results / "issues_and_assumptions.md").read_text(encoding="utf-8").lower()
    assert (results / "data_quality_report.json").is_file()
    assert (results / "issues_and_assumptions.md").is_file()
    assert (logs / "pipeline.log").is_file()
    pipeline_log = (logs / "pipeline.log").read_text(encoding="utf-8")
    assert "OPEX: Q·Σ(s_i·p_i)+ΣF_j" in pipeline_log
    assert "SPECO: C_F+Z=C_P+C_loss" in pipeline_log
    assert "Компрессоры: W=m1·w1+m2·w2" in pipeline_log
    with (logs / "harness_log.csv").open(encoding="utf-8-sig", newline="") as stream:
        rows = list(csv.DictReader(stream))
    assert any(r["Metric"] == "F-rule electricity" and r["Status"] == "PASS" for r in rows)
    book = load_workbook(results / "EA_Practical_2_Report.xlsx", read_only=True)
    assert set(book.sheetnames) >= {"Сравнение OPEX", "Энергетика и материалы", "Эксергия и SPECO", "Чувствительность"}
    assert book["Сравнение OPEX"]["A1"].value.startswith("Сравнение")
    book.close()
    document = Document(results / "EA_Practical_2_Report.docx")
    text = "\n".join(p.text for p in document.paragraphs)
    headings = [p.text for p in document.paragraphs if p.style.name.startswith("Heading")]
    assert "12. Итоговые выводы" in text
    assert "10.1. Допущения модели" in headings
    assert "10.2. Data Quality Report" in headings
    assert "7. Harness Log" in text
    assert "[ВСТАВИТЬ" not in text
    assert "TAC(ΔT)" in text
    assert "f_k = Z_k" in text
    assert "c_P,N" in text
    table_text = "\n".join(c.text for table in document.tables for row in table.rows for c in row.cells)
    assert "Аммиак, база: переменные" in table_text
    assert "ХимМаш, недогруз: доли" in table_text
