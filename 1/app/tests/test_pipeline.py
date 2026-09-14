from __future__ import annotations

import csv

import pytest
from docx import Document
from openpyxl import load_workbook

from euv_analysis.excel import ExcelRecalculator, RecalculationResult
from euv_analysis.pipeline import AnalysisPipeline


def test_pipeline_creates_docx_and_chart_without_excel_com(
    tmp_path, data_dir, monkeypatch
):
    """Catches pipeline coupling of required reports to external Excel COM."""
    monkeypatch.setattr(
        ExcelRecalculator,
        "recalculate",
        lambda self, path: RecalculationResult(
            False,
            ("EXCEL_RECALC_STARTED",),
            "COM unavailable in unit test",
        ),
    )

    result = AnalysisPipeline(data_dir=data_dir, work_dir=tmp_path).run()

    names = {path.name for path in result.output_files}
    assert "сравнение_эффективности.png" in names
    assert "Итоговый_отчет.docx" in names
    document = Document(tmp_path / "output" / "Итоговый_отчет.docx")
    assert len(document.inline_shapes) == 2
    log = (tmp_path / "logs" / "execution.log").read_text(encoding="utf-8")
    assert "EFFICIENCY_CHART_CREATED" in log
    assert "WORD_REPORT_CREATED" in log


@pytest.mark.excel
def test_pipeline_creates_verified_required_artifacts(tmp_path, data_dir):
    """Catches an incomplete pipeline or report assembled from unverified outputs."""
    result = AnalysisPipeline(data_dir=data_dir, work_dir=tmp_path).run()

    required = {
        "ground_truth.xlsx",
        "harness_log.csv",
        "stress_comparison.csv",
        "opex_structure.png",
        "сравнение_эффективности.png",
        "report.md",
        "Итоговый_отчет.docx",
    }
    assert required <= {path.name for path in result.output_files}
    assert all(path.is_file() and path.stat().st_size > 0 for path in result.output_files)
    assert (tmp_path / "logs" / "execution.log").is_file()
    assert result.excel_recalculation_success is True
    assert result.comparison_valid is True
    assert all(row.status != "FAIL" for row in result.harness_rows)
    assert all(result.boundary_checks.values())

    report = (tmp_path / "output" / "report.md").read_text(encoding="utf-8")
    assert "# Практическое занятие № 1" in report
    assert "Data Quality Report" in report
    assert "Mass Intensity" in report
    assert "NOT_COMPUTABLE" in report
    assert "TTM: сценарий задания" in report
    assert "Итоговый checklist" in report

    with (tmp_path / "output" / "stress_comparison.csv").open(
        encoding="utf-8-sig", newline=""
    ) as stream:
        stress_rows = list(csv.DictReader(stream))
    assert {row["Metric"] for row in stress_rows} == {
        "FPY",
        "OEE",
        "TEEP",
        "Energy Cost",
        "Economic Intensity",
        "CPU",
        "TTM Penalty",
    }

    workbook_path = tmp_path / "output" / "ground_truth.xlsx"
    formula_book = load_workbook(workbook_path, data_only=False, read_only=True)
    try:
        assert formula_book["Метрики"]["D2"].value.startswith("=")
        assert formula_book["Журнал сверки"].max_row > 1
    finally:
        formula_book.close()

    execution_log = (tmp_path / "logs" / "execution.log").read_text(
        encoding="utf-8"
    )
    for event in (
        "PIPELINE_STARTED",
        "CSV_READ",
        "VALIDATION_COMPLETED",
        "PYTHON_METRICS_CALCULATED",
        "EXCEL_FORMULA_WRITTEN",
        "EXCEL_RECALC_STARTED",
        "EXCEL_RECALC_COMPLETED",
        "EXCEL_VALUE_READ",
        "HARNESS_COMPLETED",
        "STRESS_TEST_COMPLETED",
        "OPEX_CHART_CREATED",
        "EFFICIENCY_CHART_CREATED",
        "REPORT_CREATED",
        "WORD_REPORT_CREATED",
        "PIPELINE_COMPLETED",
    ):
        assert event in execution_log
