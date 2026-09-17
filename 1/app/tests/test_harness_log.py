from __future__ import annotations

import json

import pytest
from openpyxl import load_workbook

from euv_analysis.harness_log import (
    HarnessLogDataError,
    HarnessLogLoader,
    HarnessLogWorkbookBuilder,
)


VALID_TEXT_FIELDS = {
    "stage": "Этап",
    "task": "Задача",
    "code_fragment": "result = calculate()",
    "issue": "Обнаруженная проблема",
    "cause": "Подтверждённая причина",
    "resolution": "Внесённое исправление",
    "recheck_result": "Повторная проверка пройдена",
}


def test_harness_log_loads_seven_verified_project_issues(data_dir):
    rows = HarnessLogLoader().load(data_dir / "harness_log.json")

    assert [row.number for row in rows] == list(range(1, 8))
    assert [row.stage for row in rows] == [
        "Аудит показателей качества",
        "Аудит параметров TTM",
        "Стресс-сценарий TTM",
        "Аудит ресурсоёмкости",
        "Пересчёт Excel",
        "Жизненный цикл COM",
        "Архитектура контрольных журналов",
    ]
    assert rows[0].recheck_result == "Тест подтверждает FPY 75 %, Final Yield 85 %"
    assert "RPC_E_DISCONNECTED" in rows[5].issue
    assert rows[6].resolution.startswith("Числовая проверка переименована")


@pytest.mark.parametrize(
    ("payload", "message"),
    [
        ([{"number": 1}], "missing fields"),
        ([{"number": 2, **VALID_TEXT_FIELDS}], "sequence"),
        ([{"number": 1, **VALID_TEXT_FIELDS, "extra": "x"}], "unknown fields"),
        ([{"number": 1, **{**VALID_TEXT_FIELDS, "stage": ""}}], "non-empty"),
    ],
)
def test_harness_log_rejects_invalid_schema(tmp_path, payload, message):
    path = tmp_path / "harness.json"
    path.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")

    with pytest.raises(HarnessLogDataError, match=message):
        HarnessLogLoader().load(path)


def test_harness_log_workbook_matches_approved_example(tmp_path, data_dir):
    entries = HarnessLogLoader().load(data_dir / "harness_log.json")
    path = HarnessLogWorkbookBuilder().build(tmp_path / "Harness_Log.xlsx", entries)
    workbook = load_workbook(path)
    sheet = workbook["Harness Log"]

    assert [cell.value for cell in sheet[1]] == [
        "№",
        "Этап",
        "Запрос к ИИ / задача",
        "Сгенерированный или изменённый фрагмент кода",
        "Обнаруженная ошибка / галлюцинация",
        "Причина",
        "Как исправлено",
        "Результат повторной проверки",
    ]
    assert sheet.max_row == 8
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref == "A1:H8"
    assert [sheet.column_dimensions[column].width for column in "ABCDEFGH"] == [
        6.0,
        24.0,
        32.0,
        40.0,
        40.0,
        35.0,
        42.0,
        38.0,
    ]
    assert all(cell.fill.fgColor.rgb == "001F4E78" for cell in sheet[1])
    assert all(cell.alignment.wrap_text for row in sheet.iter_rows() for cell in row)
