from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document

from euv_analysis.harness import MetricsHarness
from euv_analysis.loader import CsvProductionLoader
from euv_analysis.metrics import OperationalMetricsCalculator
from euv_analysis.reporting import ReportContext
from euv_analysis.stress import ComparisonCsvValidator, StressAnalyzer
from euv_analysis.ttm import TTMCalculator
from euv_analysis.validation import DataValidator
from euv_analysis.visualization import (
    EfficiencyComparisonVisualizer,
    OpexVisualizer,
)
from euv_analysis.word_reporting import WordReportBuilder


@pytest.fixture
def word_report_inputs(tmp_path: Path, data_dir: Path):
    loader = CsvProductionLoader()
    baseline = loader.load(data_dir / "euv_photoresist_pilot_2026.csv")
    stress = loader.load(data_dir / "euv_photoresist_stress_2026.csv")
    comparison = ComparisonCsvValidator().validate(
        data_dir / "euv_photoresist_comparison_input_2026.csv"
    )
    calculator = OperationalMetricsCalculator()
    baseline_metrics = calculator.calculate(baseline)
    stress_metrics = calculator.calculate(stress)
    ttm = TTMCalculator()
    assignment_ttm = ttm.calculate(
        baseline.delay_years, 0.10, 0.15, 3.0, baseline.market_window_open
    )
    csv_ttm = ttm.calculate(
        baseline.delay_years,
        baseline.discount_rate,
        baseline.price_erosion_rate,
        baseline.market_horizon,
        baseline.market_window_open,
    )
    stress_ttm = ttm.calculate(
        comparison.delay[1], 0.10, 0.15, 3.0, stress.market_window_open
    )
    python_values = baseline_metrics.numeric_values()
    python_values.update(
        {
            "TTM Penalty Assignment": assignment_ttm,
            "TTM Penalty CSV": csv_ttm,
        }
    )
    context = ReportContext(
        baseline=baseline,
        stress=stress,
        baseline_metrics=baseline_metrics,
        stress_metrics=stress_metrics,
        excel_values=python_values,
        harness_rows=MetricsHarness().compare(python_values, python_values),
        stress_rows=StressAnalyzer().compare(
            baseline_metrics,
            stress_metrics,
            baseline_ttm=0.0,
            stress_ttm=stress_ttm,
        ),
        comparison=comparison,
        validation_issues=[
            *DataValidator().validate(baseline),
            *DataValidator().validate(stress),
        ],
        boundary_checks={"Проверка полного брака": True},
        assignment_ttm=assignment_ttm,
        csv_ttm=csv_ttm,
        stress_ttm=stress_ttm,
        excel_recalculation_success=True,
    )
    opex_path = OpexVisualizer().create(tmp_path / "opex.png")
    efficiency_path = EfficiencyComparisonVisualizer().create(
        tmp_path / "efficiency.png", baseline_metrics, stress_metrics
    )
    return context, opex_path, efficiency_path


def test_word_report_contains_calculated_tables_and_two_charts(
    tmp_path: Path, word_report_inputs
):
    """Catches a DOCX with placeholders, hard-coded results, or missing charts."""
    context, opex_path, efficiency_path = word_report_inputs
    path = tmp_path / "Итоговый_отчет.docx"

    result = WordReportBuilder().build(
        path,
        context,
        opex_chart_path=opex_path,
        efficiency_chart_path=efficiency_path,
    )

    document = Document(result)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    table_text = "\n".join(
        cell.text
        for table in document.tables
        for row in table.rows
        for cell in row.cells
    )
    assert result.stat().st_size > 50_000
    assert "Итоговый отчёт по экономическому анализу" in text
    assert "Результаты расчёта" in text
    assert "Стресс-сценарий" in text
    assert "26,2859 %" in text
    assert "НЕ РАССЧИТЫВАЕТСЯ" in table_text
    assert "14 713,04" in table_text
    assert len(document.tables) >= 6
    assert len(document.inline_shapes) == 2


def test_word_report_rejects_missing_chart(tmp_path: Path, word_report_inputs):
    """Catches silently publishing a Word report without required evidence."""
    context, opex_path, _ = word_report_inputs
    missing = tmp_path / "missing.png"

    with pytest.raises(FileNotFoundError, match="missing.png"):
        WordReportBuilder().build(
            tmp_path / "report.docx",
            context,
            opex_chart_path=opex_path,
            efficiency_chart_path=missing,
        )

