from __future__ import annotations

from pathlib import Path
import os
import subprocess
import sys

import pytest
from openpyxl import load_workbook

from euv_analysis.excel import (
    ExcelGroundTruthBuilder,
    ExcelGroundTruthReader,
    ExcelRecalculator,
)
from euv_analysis.harness import HarnessEntry
from euv_analysis.loader import CsvProductionLoader


def test_workbook_contains_independent_formulas(tmp_path, valid_data):
    """Catches a workbook that copies Python results instead of calculating them."""
    path = ExcelGroundTruthBuilder().build(
        tmp_path / "ground_truth.xlsx", valid_data, valid_data
    )

    workbook = load_workbook(path, data_only=False)
    assert {
        "Исходные данные",
        "Метрики",
        "Затраты",
        "TTM",
        "Журнал сверки",
        "Стресс-тест",
        "Структура OPEX",
        "Допущения",
        "Проблемы данных",
    } <= set(workbook.sheetnames)
    formulas = [
        cell.value
        for sheet_name in ("Метрики", "Затраты", "TTM", "Стресс-тест")
        for row in workbook[sheet_name].iter_rows()
        for cell in row
        if cell.data_type == "f"
    ]
    assert len(formulas) >= 20
    assert any("'Исходные данные'!" in formula for formula in formulas)
    assert workbook["Метрики"]["D2"].value.startswith("=")
    assert workbook["Метрики"]["D2"].value != "=0.75"


def test_workbook_uses_russian_sheet_names_headers_and_values(tmp_path, valid_data):
    """Catches regression to English labels in the user-facing workbook."""
    path = ExcelGroundTruthBuilder().build(
        tmp_path / "ground_truth.xlsx", valid_data, valid_data
    )
    workbook = load_workbook(path, data_only=False)

    assert workbook.sheetnames == [
        "Исходные данные",
        "Метрики",
        "Затраты",
        "TTM",
        "Стресс-тест",
        "Журнал сверки",
        "Структура OPEX",
        "Допущения",
        "Проблемы данных",
    ]
    assert [cell.value for cell in workbook["Исходные данные"][1]] == [
        "Сценарий",
        "Параметр",
        "Техническое имя CSV",
        "Значение",
        "Единица измерения",
        "Описание",
    ]
    assert workbook["Исходные данные"]["A2"].value == "Базовый"
    assert workbook["Исходные данные"]["B2"].value == "Календарное время"
    assert workbook["Исходные данные"]["C2"].value == "Calendar_Hours"
    assert (
        workbook["Метрики"]["A2"].value
        == "Выход годных с первого прохода (FPY)"
    )
    assert (
        workbook["Метрики"]["B2"].value
        == "Годный объём первого прохода / фактический выпуск"
    )
    assert workbook["Метрики"]["F2"].value.startswith("Доля продукции")
    assert workbook["Метрики"]["D13"].value == "НЕ РАССЧИТЫВАЕТСЯ"


def test_workbook_records_mass_intensity_data_gap(tmp_path, valid_data):
    """Catches a workbook that invents missing raw-material mass."""
    path = ExcelGroundTruthBuilder().build(
        tmp_path / "ground_truth.xlsx", valid_data, valid_data
    )
    workbook = load_workbook(path, data_only=False)

    mass_row = next(
        row
        for row in workbook["Метрики"].iter_rows(values_only=True)
        if row[0] == "Массовая ресурсоёмкость"
    )
    assert mass_row[3] == "НЕ РАССЧИТЫВАЕТСЯ"
    assert "Raw_Material_Mass_kg" in mass_row[5]


def test_workbook_stress_sheet_includes_ttm_penalty_formulas(tmp_path, valid_data):
    """Catches omission of the required TTM metric from Excel stress results."""
    path = ExcelGroundTruthBuilder().build(
        tmp_path / "ground_truth.xlsx", valid_data, valid_data
    )
    workbook = load_workbook(path, data_only=False)

    ttm_row = next(
        row
        for row in workbook["Стресс-тест"].iter_rows(values_only=True)
        if row[0] == "Штраф за задержку вывода на рынок (TTM)"
    )
    assert ttm_row[1].startswith("=")
    assert ttm_row[2].startswith("=")


def test_harness_rows_are_localized_only_in_workbook(tmp_path, valid_data):
    """Catches leaking internal English metric names and statuses into Excel."""
    path = ExcelGroundTruthBuilder().build(
        tmp_path / "ground_truth.xlsx", valid_data, valid_data
    )
    entry = HarnessEntry(
        "OEE",
        0.5,
        0.5,
        0.0,
        0.0,
        1e-9,
        "PASS",
        "",
        "2026-09-14T00:00:00+00:00",
    )

    ExcelGroundTruthBuilder().write_harness(path, [entry])

    workbook = load_workbook(path, data_only=False)
    row = list(
        workbook["Журнал сверки"].iter_rows(min_row=2, values_only=True)
    )[0]
    assert row[0] == "Общая эффективность оборудования (OEE)"
    assert row[6] == "СОВПАДАЕТ"


class _FakeWorkbook:
    def __init__(self):
        self.saved = False
        self.closed = False

    def Save(self):
        self.saved = True

    def Close(self, SaveChanges=False):
        self.closed = True


class _FakeWorkbooks:
    def __init__(self, workbook, fail=False):
        self.workbook = workbook
        self.fail = fail
        self.opened_path = None

    def Open(self, path):
        if self.fail:
            raise RuntimeError("open failed")
        self.opened_path = path
        return self.workbook


class _FakeExcel:
    def __init__(self, fail=False):
        self.workbook = _FakeWorkbook()
        self.Workbooks = _FakeWorkbooks(self.workbook, fail=fail)
        self.DisplayAlerts = True
        self.Visible = True
        self.calculated = False
        self.quit_called = False

    def CalculateFullRebuild(self):
        self.calculated = True

    def Quit(self):
        self.quit_called = True


def test_excel_recalculator_reports_ordered_success_events(tmp_path):
    """Catches a recalculator that reports success before save and close."""
    path = tmp_path / "book.xlsx"
    path.touch()
    fake = _FakeExcel()

    result = ExcelRecalculator(excel_factory=lambda: fake).recalculate(path)

    assert result.success is True
    assert result.events == (
        "EXCEL_RECALC_STARTED",
        "EXCEL_RECALC_COMPLETED",
    )
    assert fake.Workbooks.opened_path == str(path.resolve())
    assert fake.calculated is True
    assert fake.workbook.saved is True
    assert fake.workbook.closed is True
    assert fake.quit_called is True


def test_excel_recalculator_does_not_mask_failure(tmp_path):
    """Catches a false successful Ground Truth when Excel cannot open the file."""
    path = tmp_path / "book.xlsx"
    path.touch()
    fake = _FakeExcel(fail=True)

    result = ExcelRecalculator(excel_factory=lambda: fake).recalculate(path)

    assert result.success is False
    assert "open failed" in result.error
    assert "EXCEL_RECALC_COMPLETED" not in result.events
    assert fake.quit_called is True


@pytest.mark.excel
def test_installed_excel_recalculates_formula_cache(tmp_path, data_dir):
    """Catches COM integration that saves formulas without cached results."""
    loader = CsvProductionLoader()
    baseline = loader.load(data_dir / "euv_photoresist_pilot_2026.csv")
    stress = loader.load(data_dir / "euv_photoresist_stress_2026.csv")
    path = ExcelGroundTruthBuilder().build(
        tmp_path / "ground_truth.xlsx", baseline, stress
    )

    recalculation = ExcelRecalculator().recalculate(path)
    values = ExcelGroundTruthReader().read(path)

    assert recalculation.success, recalculation.error
    assert values["FPY"] == pytest.approx(0.75)
    assert values["OEE"] == pytest.approx(0.59375)
    assert values["CPU"] == pytest.approx(14_713.037495700035)


@pytest.mark.excel
def test_excel_recalculation_releases_com_without_rpc_fault(tmp_path, data_dir):
    """Catches COM proxies surviving Excel.Quit and faulting during teardown."""
    loader = CsvProductionLoader()
    path = ExcelGroundTruthBuilder().build(
        tmp_path / "ground_truth.xlsx",
        loader.load(data_dir / "euv_photoresist_pilot_2026.csv"),
        loader.load(data_dir / "euv_photoresist_stress_2026.csv"),
    )
    app_dir = Path(__file__).resolve().parents[1]
    environment = os.environ.copy()
    environment["PYTHONPATH"] = str(app_dir / "src")
    completed = subprocess.run(
        [
            sys.executable,
            "-X",
            "faulthandler",
            "-c",
            (
                "from pathlib import Path; "
                "from euv_analysis.excel import ExcelRecalculator, ExcelGroundTruthReader; "
                f"result=ExcelRecalculator().recalculate(Path({str(path)!r})); "
                f"value=ExcelGroundTruthReader().read(Path({str(path)!r}))['FPY']; "
                "raise SystemExit(0 if result.success and value == 0.75 else 1)"
            ),
        ],
        cwd=app_dir,
        env=environment,
        capture_output=True,
        text=True,
        timeout=30,
        check=False,
    )

    assert completed.returncode == 0, completed.stderr
    assert "Windows fatal exception" not in completed.stderr
    assert "0x80010108" not in completed.stderr
