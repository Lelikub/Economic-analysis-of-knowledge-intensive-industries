"""Final Word report assembled from already calculated pipeline evidence."""

from __future__ import annotations

from pathlib import Path
from collections.abc import Iterable, Sequence

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from .excel import PARAMETER_FIELDS
from .localization import (
    PARAMETER_LABELS,
    localize_metric,
    localize_status,
    localize_unit,
)
from .reporting import ReportContext


_FORMULA_EXPLANATIONS: dict[str, str] = {
    "FPY": "Годный объём первого прохода / фактический выпуск",
    "Final Yield": "Итоговый годный объём / фактический выпуск",
    "Availability": "Фактическое время работы / плановое время работы",
    "Performance": "Нормативное время цикла × фактический выпуск / фактическое время",
    "Quality": "Коэффициент качества равен FPY",
    "OEE": "Доступность × производительность × качество",
    "Utilization": "Плановое время работы / календарное время",
    "TEEP": "OEE × использование календарного времени",
    "Energy Cost": "Потребление электроэнергии × тариф",
    "Monthly Depreciation": "CAPEX / экономический срок службы / 12",
    "Total Manufacturing Cost": "Сырьё + энергия + накладные расходы + амортизация",
    "Mass Intensity": "Масса сырья / итоговый годный объём",
    "Energy Intensity": "Потребление электроэнергии / итоговый годный объём",
    "Economic Intensity": "Общие производственные затраты / итоговый годный объём",
    "Raw Material Cost Proxy": "Затраты на сырьё / итоговый годный объём",
    "CPU": "Общие производственные затраты / итоговый годный объём",
}


def _number(value: float, digits: int = 2) -> str:
    return f"{value:,.{digits}f}".replace(",", " ").replace(".", ",")


def _percent(value: float, digits: int = 2) -> str:
    return f"{_number(value * 100, digits)} %"


def _metric_value(value: float | None, unit: str) -> str:
    if value is None:
        return localize_status("NOT_COMPUTABLE")
    if unit in {"dimensionless", "ratio"}:
        return _percent(value)
    return _number(value)


def _parameter_value(value: float | int, unit: str) -> str:
    if unit == "ratio":
        return _percent(float(value))
    if unit == "binary":
        return "Открыто" if int(value) == 1 else "Закрыто"
    return _number(float(value), 4).rstrip("0").rstrip(",")


def _set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def _add_table(
    document: Document,
    headers: Sequence[str],
    rows: Iterable[Sequence[object]],
):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = header
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        _set_cell_shading(cell, "1F4E78")
        for run in cell.paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
    for values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].text = str(value)
            cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    document.add_paragraph()
    return table


class WordReportBuilder:
    """Build a styled DOCX without recalculating any business metric."""

    def build(
        self,
        path: Path,
        context: ReportContext,
        *,
        opex_chart_path: Path,
        efficiency_chart_path: Path,
    ) -> Path:
        """Create the final report from a shared calculated context."""
        path = Path(path)
        opex_chart_path = Path(opex_chart_path)
        efficiency_chart_path = Path(efficiency_chart_path)
        for chart_path in (opex_chart_path, efficiency_chart_path):
            if not chart_path.is_file():
                raise FileNotFoundError(f"Не найден обязательный график: {chart_path}")

        document = Document()
        self._configure_document(document)
        self._add_title(document)
        self._add_scope(document)
        self._add_parameters(document, context)
        self._add_formulas(document, context)
        self._add_metrics(document, context)
        self._add_harness(document, context)
        self._add_stress(document, context)
        self._add_ttm(document, context)
        self._add_charts(document, opex_chart_path, efficiency_chart_path)
        self._add_quality_and_assumptions(document, context)
        self._add_conclusions(document, context)

        path.parent.mkdir(parents=True, exist_ok=True)
        temporary_path = path.with_name(f".{path.name}.tmp")
        document.save(temporary_path)
        temporary_path.replace(path)
        return path

    @staticmethod
    def _configure_document(document: Document) -> None:
        section = document.sections[0]
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(1.8)
        section.bottom_margin = Cm(1.8)
        section.left_margin = Cm(1.8)
        section.right_margin = Cm(1.5)

        normal = document.styles["Normal"]
        normal.font.name = "Times New Roman"
        normal.font.size = Pt(10.5)
        normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        for style_name, size, color in (
            ("Title", 20, "1F4E78"),
            ("Heading 1", 15, "1F4E78"),
            ("Heading 2", 12, "2F75B5"),
        ):
            style = document.styles[style_name]
            style.font.name = "Times New Roman"
            style.font.size = Pt(size)
            style.font.color.rgb = RGBColor.from_string(color)
            style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

    @staticmethod
    def _add_title(document: Document) -> None:
        title = document.add_heading(
            "Итоговый отчёт по экономическому анализу", level=0
        )
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        subtitle = document.add_paragraph(
            "Пилотное производство EUV-фоторезиста\nПрактическое занятие № 1"
        )
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
        document.add_paragraph()

    @staticmethod
    def _add_scope(document: Document) -> None:
        document.add_heading("1. Цель и исходные данные", level=1)
        document.add_paragraph(
            "Цель работы — связать физические параметры производства с "
            "эффективностью оборудования, ресурсоёмкостью, себестоимостью и "
            "экономическим штрафом за задержку вывода продукта на рынок."
        )
        document.add_paragraph(
            "Использованы файлы euv_photoresist_pilot_2026.csv, "
            "euv_photoresist_stress_2026.csv и "
            "euv_photoresist_comparison_input_2026.csv."
        )

    @staticmethod
    def _add_parameters(document: Document, context: ReportContext) -> None:
        document.add_heading("2. Параметры модели", level=1)
        rows = []
        for source_name, field_name in PARAMETER_FIELDS:
            unit = context.baseline.units.get(source_name, "")
            rows.append(
                (
                    PARAMETER_LABELS[source_name],
                    source_name,
                    _parameter_value(getattr(context.baseline, field_name), unit),
                    _parameter_value(getattr(context.stress, field_name), unit),
                    localize_unit(unit),
                )
            )
        _add_table(
            document,
            (
                "Параметр",
                "Техническое имя CSV",
                "Базовый сценарий",
                "Стрессовый сценарий",
                "Единица измерения",
            ),
            rows,
        )

    @staticmethod
    def _add_formulas(document: Document, context: ReportContext) -> None:
        document.add_heading("3. Формулы и пояснения", level=1)
        rows = []
        for field_name in context.baseline_metrics.__dataclass_fields__:
            metric = getattr(context.baseline_metrics, field_name)
            rows.append(
                (
                    localize_metric(metric.name),
                    _FORMULA_EXPLANATIONS[metric.name],
                    localize_unit(metric.unit),
                )
            )
        _add_table(
            document,
            ("Метрика", "Пояснение формулы", "Единица измерения"),
            rows,
        )

    @staticmethod
    def _add_metrics(document: Document, context: ReportContext) -> None:
        document.add_heading("4. Результаты расчёта", level=1)
        rows = []
        for field_name in context.baseline_metrics.__dataclass_fields__:
            baseline = getattr(context.baseline_metrics, field_name)
            stress = getattr(context.stress_metrics, field_name)
            rows.append(
                (
                    localize_metric(baseline.name),
                    _metric_value(baseline.value, baseline.unit),
                    _metric_value(stress.value, stress.unit),
                    localize_unit(baseline.unit),
                    localize_status(baseline.status.value),
                )
            )
        _add_table(
            document,
            (
                "Метрика",
                "Базовый сценарий",
                "Стрессовый сценарий",
                "Единица измерения",
                "Статус",
            ),
            rows,
        )

    @staticmethod
    def _add_harness(document: Document, context: ReportContext) -> None:
        document.add_heading("5. Сверка Python и Excel", level=1)
        status = "выполнен" if context.excel_recalculation_success else "недоступен"
        document.add_paragraph(f"Реальный пересчёт Microsoft Excel: {status}.")
        rows = [
            (
                localize_metric(row.metric),
                _number(row.excel_value, 8) if row.excel_value is not None else "Нет значения",
                _number(row.python_value, 8) if row.python_value is not None else "Нет значения",
                _number(row.absolute_delta, 10) if row.absolute_delta is not None else "—",
                localize_status(row.status),
            )
            for row in context.harness_rows
        ]
        _add_table(
            document,
            ("Метрика", "Excel", "Python", "Абсолютное отклонение", "Статус"),
            rows,
        )

    @staticmethod
    def _add_stress(document: Document, context: ReportContext) -> None:
        document.add_heading("6. Стресс-сценарий", level=1)
        percent_metrics = {"FPY", "OEE", "TEEP", "TTM Penalty"}
        rows = []
        for row in context.stress_rows:
            formatter = _percent if row.metric in percent_metrics else _number
            rows.append(
                (
                    localize_metric(row.metric),
                    formatter(row.baseline) if row.baseline is not None else "Нет значения",
                    formatter(row.stress) if row.stress is not None else "Нет значения",
                    formatter(row.absolute_change) if row.absolute_change is not None else "Нет значения",
                    _percent(row.relative_change) if row.relative_change is not None else "Не определяется",
                    row.interpretation,
                )
            )
        _add_table(
            document,
            (
                "Метрика",
                "Базовый",
                "Стрессовый",
                "Абсолютное изменение",
                "Относительное изменение",
                "Интерпретация",
            ),
            rows,
        )

    @staticmethod
    def _add_ttm(document: Document, context: ReportContext) -> None:
        document.add_heading("7. Время вывода на рынок (TTM)", level=1)
        _add_table(
            document,
            (
                "Сценарий",
                "Ставка дисконтирования",
                "Ценовая эрозия",
                "Горизонт",
                "Задержка",
                "Штраф TTM",
            ),
            (
                ("Параметры задания", "10,00 %", "15,00 %", "3 года", "0,5 года", _percent(context.assignment_ttm, 4)),
                ("Параметры CSV", "12,00 %", "18,00 %", "3 года", "0,5 года", _percent(context.csv_ttm, 4)),
                ("Стресс относительно базового сравнения", "10,00 %", "15,00 %", "3 года", "0,5 года", _percent(context.stress_ttm, 4)),
            ),
        )
        document.add_paragraph(
            "Штраф TTM по параметрам задания составляет "
            f"{_percent(context.assignment_ttm, 4)}. Он учитывает дисконтирование, "
            "ценовую эрозию и сокращение доступного рыночного горизонта."
        )

    @staticmethod
    def _add_charts(
        document: Document, opex_chart_path: Path, efficiency_chart_path: Path
    ) -> None:
        document.add_heading("8. Графики", level=1)
        document.add_heading("8.1. Структура OPEX", level=2)
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run().add_picture(str(opex_chart_path), width=Cm(16))
        document.add_paragraph(
            "Рисунок 1 — Структура операционных расходов высокотехнологичного производства"
        ).alignment = WD_ALIGN_PARAGRAPH.CENTER
        document.add_heading("8.2. Сравнение эффективности", level=2)
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run().add_picture(str(efficiency_chart_path), width=Cm(16))
        document.add_paragraph(
            "Рисунок 2 — FPY, OEE и TEEP в базовом и стрессовом сценариях"
        ).alignment = WD_ALIGN_PARAGRAPH.CENTER

    @staticmethod
    def _add_quality_and_assumptions(
        document: Document, context: ReportContext
    ) -> None:
        document.add_heading("9. Допущения и качество данных", level=1)
        _add_table(
            document,
            ("Код", "Допущение", "Обоснование", "Влияние"),
            (
                ("A01", "Расчётный период — один месяц", "720 часов = 30 × 24", "Ежемесячная амортизация"),
                ("A02", "Нормативное время цикла получено из плана", "Отдельного поля нет", "Производительность и OEE"),
                ("A03", "В знаменателе CPU используется итоговый годный объём", "Это объём после итогового контроля", "CPU и ресурсоёмкость"),
            ),
        )
        issue_rows = [
            ("Отсутствует масса сырья", "Raw_Material_Mass_kg", "ВЫСОКАЯ", "Массовая ресурсоёмкость не рассчитывается", "Добавить массу сырья в килограммах"),
            ("Параметры TTM расходятся", "Discount_Rate; Price_Erosion_Rate", "СРЕДНЯЯ", "Получаются разные штрафы", "Показывать сценарии отдельно"),
            ("Задержка пилотного сценария отличается от базового сравнения", "Delay_Years", "СРЕДНЯЯ", "Пилот не является сценарием без задержки", "Использовать сравнительный CSV"),
        ]
        issue_rows.extend(
            (
                issue.message,
                issue.parameter,
                issue.severity.value,
                issue.category,
                "Исправить исходные данные или подтвердить стрессовый режим",
            )
            for issue in context.validation_issues
        )
        _add_table(
            document,
            ("Проблема", "Параметр", "Критичность", "Влияние", "Действие"),
            issue_rows,
        )

    @staticmethod
    def _add_conclusions(document: Document, context: ReportContext) -> None:
        document.add_heading("10. Выводы", level=1)
        factors = {
            "доступность": context.baseline_metrics.availability.value,
            "производительность": context.baseline_metrics.performance.value,
            "качество (FPY)": context.baseline_metrics.fpy.value,
        }
        limiting_factor = min(
            (item for item in factors.items() if item[1] is not None),
            key=lambda item: float(item[1]),
        )[0]
        stress_index = {row.metric: row for row in context.stress_rows}
        document.add_paragraph(
            "Главным ограничивающим фактором OEE является "
            f"{limiting_factor}. В базовом сценарии FPY равен "
            f"{_percent(context.baseline_metrics.fpy.value or 0.0)}, OEE — "
            f"{_percent(context.baseline_metrics.oee.value or 0.0)}, а TEEP — "
            f"{_percent(context.baseline_metrics.teep.value or 0.0)}."
        )
        document.add_paragraph(
            "В стрессовом сценарии FPY снижается до "
            f"{_percent(context.stress_metrics.fpy.value or 0.0)}, OEE — до "
            f"{_percent(context.stress_metrics.oee.value or 0.0)}, TEEP — до "
            f"{_percent(context.stress_metrics.teep.value or 0.0)}."
        )
        cpu_change = stress_index["CPU"].absolute_change
        energy_change = stress_index["Energy Cost"].absolute_change
        document.add_paragraph(
            "Трёхкратное повышение тарифа увеличивает затраты на электроэнергию "
            f"на {_number(energy_change or 0.0)} долл. США и повышает CPU на "
            f"{_number(cpu_change or 0.0)} долл. США/л. Массовая ресурсоёмкость "
            "не рассчитывается до появления физического параметра массы сырья."
        )

