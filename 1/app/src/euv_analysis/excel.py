"""Independent Excel Ground Truth workbook and recalculation adapter."""

from __future__ import annotations

from dataclasses import dataclass
import logging
from pathlib import Path
import subprocess
import sys
from typing import Any, Callable, Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.drawing.image import Image
from openpyxl.styles import Alignment, Font, PatternFill

from .localization import (
    PARAMETER_DESCRIPTIONS,
    PARAMETER_LABELS,
    SCENARIO_LABELS,
    SHEET_NAMES,
    localize_cause,
    localize_metric,
    localize_status,
    localize_unit,
)
from .models import ProductionData


LOGGER = logging.getLogger(__name__)


PARAMETER_FIELDS: tuple[tuple[str, str], ...] = (
    ("Calendar_Hours", "calendar_hours"),
    ("Planned_Hours", "planned_hours"),
    ("Actual_Hours", "actual_hours"),
    ("Target_Output", "target_output"),
    ("Actual_Output", "actual_output"),
    ("First_Pass_Good_Units", "first_pass_good_units"),
    ("Final_Good_Units", "final_good_units"),
    ("Raw_Material_Cost", "raw_material_cost"),
    ("Energy_kWh", "energy_kwh"),
    ("Energy_Cost_Rate", "energy_cost_rate"),
    ("Equipment_CAPEX", "equipment_capex"),
    ("Amortization_Years_Economic", "amortization_years_economic"),
    ("OPEX_Overhead", "opex_overhead"),
    ("Delay_Years", "delay_years"),
    ("Discount_Rate", "discount_rate"),
    ("Price_Erosion_Rate", "price_erosion_rate"),
    ("Market_Horizon", "market_horizon"),
    ("Market_Window_Open", "market_window_open"),
)


GROUND_TRUTH_CELLS: dict[str, tuple[str, str]] = {
    "FPY": (SHEET_NAMES["metrics"], "D2"),
    "Final Yield": (SHEET_NAMES["metrics"], "D3"),
    "Availability": (SHEET_NAMES["metrics"], "D4"),
    "Performance": (SHEET_NAMES["metrics"], "D5"),
    "Quality": (SHEET_NAMES["metrics"], "D6"),
    "OEE": (SHEET_NAMES["metrics"], "D7"),
    "Utilization": (SHEET_NAMES["metrics"], "D8"),
    "TEEP": (SHEET_NAMES["metrics"], "D9"),
    "Energy Intensity": (SHEET_NAMES["metrics"], "D10"),
    "Economic Intensity": (SHEET_NAMES["metrics"], "D11"),
    "CPU": (SHEET_NAMES["metrics"], "D12"),
    "Mass Intensity": (SHEET_NAMES["metrics"], "D13"),
    "TTM Penalty Assignment": (SHEET_NAMES["metrics"], "D14"),
    "TTM Penalty CSV": (SHEET_NAMES["metrics"], "D15"),
    "Energy Cost": (SHEET_NAMES["costs"], "D2"),
    "Monthly Depreciation": (SHEET_NAMES["costs"], "D3"),
    "Total Manufacturing Cost": (SHEET_NAMES["costs"], "D4"),
    "Raw Material Cost Proxy": (SHEET_NAMES["costs"], "D7"),
}


def _style_sheet(sheet, widths: dict[str, float]) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width


def _ttm_formula(row: int) -> str:
    return (
        f"=IF(OR(F{row}=0,E{row}>=D{row}),1,"
        f"1-(1/(1+B{row})^E{row})*EXP(-C{row}*E{row})*"
        f"((D{row}-E{row})/D{row})*F{row})"
    )


class ExcelGroundTruthBuilder:
    """Build an xlsx workbook whose results are native Excel formulas."""

    def build(
        self,
        path: Path,
        baseline: ProductionData,
        stress: ProductionData,
        *,
        opex_chart_path: Path | None = None,
        stress_baseline_delay: float = 0.0,
        assumptions: list[dict[str, str]] | None = None,
        data_issues: list[dict[str, str]] | None = None,
    ) -> Path:
        """Create the Ground Truth workbook without using Python metrics."""
        path = Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)
        workbook = Workbook()
        workbook.remove(workbook.active)
        workbook.calculation.calcMode = "auto"
        workbook.calculation.fullCalcOnLoad = True
        workbook.calculation.forceFullCalc = True

        source = workbook.create_sheet(SHEET_NAMES["source"])
        source.append(
            [
                "Сценарий",
                "Параметр",
                "Техническое имя CSV",
                "Значение",
                "Единица измерения",
                "Описание",
            ]
        )
        refs: dict[str, dict[str, str]] = {"Baseline": {}, "Stress": {}}
        for scenario_name, data in (("Baseline", baseline), ("Stress", stress)):
            for parameter_name, field_name in PARAMETER_FIELDS:
                row_number = source.max_row + 1
                source.append(
                    [
                        SCENARIO_LABELS[scenario_name],
                        PARAMETER_LABELS[parameter_name],
                        parameter_name,
                        getattr(data, field_name),
                        localize_unit(data.units.get(parameter_name, "")),
                        PARAMETER_DESCRIPTIONS[parameter_name],
                    ]
                )
                refs[scenario_name][parameter_name] = (
                    f"'{SHEET_NAMES['source']}'!$D${row_number}"
                )
        _style_sheet(
            source,
            {"A": 18, "B": 38, "C": 36, "D": 18, "E": 24, "F": 66},
        )

        b = refs["Baseline"]
        s = refs["Stress"]
        metrics = workbook.create_sheet(SHEET_NAMES["metrics"])
        metrics.append(
            [
                "Метрика",
                "Пояснение формулы",
                "Используемые параметры",
                "Значение",
                "Единица измерения",
                "Комментарий",
            ]
        )
        metric_rows = [
            (
                localize_metric("FPY"),
                "Годный объём первого прохода / фактический выпуск",
                "Годный объём первого прохода; фактический выпуск",
                f"={b['First_Pass_Good_Units']}/{b['Actual_Output']}",
                localize_unit("dimensionless"),
                "Доля продукции, прошедшей контроль качества без переделки",
            ),
            (
                localize_metric("Final Yield"),
                "Итоговый годный объём / фактический выпуск",
                "Итоговый годный объём; фактический выпуск",
                f"={b['Final_Good_Units']}/{b['Actual_Output']}",
                localize_unit("dimensionless"),
                "Учитывает годную продукцию после переделки",
            ),
            (
                localize_metric("Availability"),
                "Фактическое время работы / плановое время работы",
                "Фактическое время работы; плановое время работы",
                f"={b['Actual_Hours']}/{b['Planned_Hours']}",
                localize_unit("dimensionless"),
                "Коэффициент доступности в составе OEE",
            ),
            (
                localize_metric("Performance"),
                "(Плановое время / плановый выпуск) × фактический выпуск / фактическое время",
                "Плановое время; плановый выпуск; фактический выпуск; фактическое время",
                f"=({b['Planned_Hours']}/{b['Target_Output']})*{b['Actual_Output']}/{b['Actual_Hours']}",
                localize_unit("dimensionless"),
                "Нормативное время цикла рассчитано в часах на литр",
            ),
            (
                localize_metric("Quality"),
                "Коэффициент качества = FPY",
                "Выход годных с первого прохода (FPY)",
                "=D2",
                localize_unit("dimensionless"),
                "Определение коэффициента качества по лекции",
            ),
            (
                localize_metric("OEE"),
                "Доступность × производительность × качество",
                "Доступность; производительность; качество",
                "=D4*D5*D6",
                localize_unit("dimensionless"),
                "Операционная эффективность оборудования",
            ),
            (
                localize_metric("Utilization"),
                "Плановое время работы / календарное время",
                "Плановое время работы; календарное время",
                f"={b['Planned_Hours']}/{b['Calendar_Hours']}",
                localize_unit("dimensionless"),
                "Использование доступного календарного фонда времени",
            ),
            (
                localize_metric("TEEP"),
                "OEE × использование календарного времени",
                "OEE; использование календарного времени",
                "=D7*D8",
                localize_unit("dimensionless"),
                "Показывает использование производственного потенциала CAPEX",
            ),
            (
                localize_metric("Energy Intensity"),
                "Потребление электроэнергии / итоговый годный объём",
                "Потребление электроэнергии; итоговый годный объём",
                f"={b['Energy_kWh']}/{b['Final_Good_Units']}",
                localize_unit("kWh/L"),
                "Потребление энергии на литр итоговой годной продукции",
            ),
            (
                localize_metric("Economic Intensity"),
                "Общие производственные затраты / итоговый годный объём",
                "Общие производственные затраты; итоговый годный объём",
                f"='{SHEET_NAMES['costs']}'!D5",
                localize_unit("USD/L"),
                "В данной области расчёта численно совпадает с CPU",
            ),
            (
                localize_metric("CPU"),
                "Общие производственные затраты / итоговый годный объём",
                "Общие производственные затраты; итоговый годный объём",
                f"='{SHEET_NAMES['costs']}'!D6",
                localize_unit("USD/L"),
                "Себестоимость литра итоговой годной продукции",
            ),
            (
                localize_metric("Mass Intensity"),
                "Масса сырья / итоговый годный объём",
                "Raw_Material_Mass_kg (масса сырья)",
                localize_status("NOT_COMPUTABLE"),
                localize_unit("kg/L"),
                "Параметр Raw_Material_Mass_kg отсутствует в исходных данных",
            ),
            (
                localize_metric("TTM Penalty Assignment"),
                "Мультипликативная потеря сохраняемой стоимости из-за задержки",
                "Параметры TTM в строке 2",
                f"='{SHEET_NAMES['ttm']}'!G2",
                localize_unit("dimensionless"),
                "Ставка дисконтирования 10 %, ценовая эрозия 15 %, горизонт 3 года",
            ),
            (
                localize_metric("TTM Penalty CSV"),
                "Мультипликативная потеря сохраняемой стоимости из-за задержки",
                "Параметры TTM в строке 3",
                f"='{SHEET_NAMES['ttm']}'!G3",
                localize_unit("dimensionless"),
                "Ставки из CSV показаны отдельным сценарием",
            ),
        ]
        for row in metric_rows:
            metrics.append(row)
            if isinstance(row[3], str) and row[3].startswith("="):
                LOGGER.info("EXCEL_FORMULA_WRITTEN metric=%s", row[0])
        for row in range(2, metrics.max_row + 1):
            if metrics.cell(row, 5).value == localize_unit("dimensionless"):
                metrics.cell(row, 4).number_format = "0.0000%"
        _style_sheet(metrics, {"A": 28, "B": 48, "C": 58, "D": 20, "E": 18, "F": 58})

        costs = workbook.create_sheet(SHEET_NAMES["costs"])
        costs.append(
            [
                "Метрика",
                "Пояснение формулы",
                "Используемые параметры",
                "Значение",
                "Единица измерения",
                "Комментарий",
            ]
        )
        cost_rows = [
            (localize_metric("Energy Cost"), "Потребление электроэнергии × тариф", "Потребление электроэнергии; тариф на электроэнергию", f"={b['Energy_kWh']}*{b['Energy_Cost_Rate']}", localize_unit("USD/month"), "Затраты на покупную электроэнергию"),
            (localize_metric("Monthly Depreciation"), "CAPEX / экономический срок / 12", "Капитальная стоимость; экономический срок амортизации", f"={b['Equipment_CAPEX']}/{b['Amortization_Years_Economic']}/12", localize_unit("USD/month"), "Набор данных за 30 суток принят за один месяц"),
            (localize_metric("Total Manufacturing Cost"), "Сырьё + энергия + накладные расходы + амортизация", "Исходные данные; затраты на электроэнергию; амортизация", f"={b['Raw_Material_Cost']}+D2+{b['OPEX_Overhead']}+D3", localize_unit("USD/month"), "Прозрачная структура полной производственной стоимости"),
            (localize_metric("Economic Intensity"), "Общие производственные затраты / итоговый годный объём", "Общие производственные затраты; итоговый годный объём", f"=D4/{b['Final_Good_Units']}", localize_unit("USD/L"), "Полная экономическая ресурсоёмкость"),
            (localize_metric("CPU"), "Общие производственные затраты / итоговый годный объём", "Общие производственные затраты; итоговый годный объём", f"=D4/{b['Final_Good_Units']}", localize_unit("USD/L"), "Себестоимость единицы продукции"),
            (localize_metric("Raw Material Cost Proxy"), "Затраты на сырьё / итоговый годный объём", "Затраты на сырьё; итоговый годный объём", f"={b['Raw_Material_Cost']}/{b['Final_Good_Units']}", localize_unit("USD/L"), "Стоимостной показатель не заменяет массовую ресурсоёмкость"),
        ]
        for row in cost_rows:
            costs.append(row)
            LOGGER.info("EXCEL_FORMULA_WRITTEN metric=%s", row[0])
        _style_sheet(costs, {"A": 30, "B": 50, "C": 52, "D": 20, "E": 18, "F": 58})

        ttm = workbook.create_sheet(SHEET_NAMES["ttm"])
        ttm.append(["Сценарий", "Ставка дисконтирования", "Ценовая эрозия", "Рыночный горизонт, лет", "Задержка, лет", "Рыночное окно открыто", "Штраф TTM", "Комментарий"])
        ttm.append([SCENARIO_LABELS["Assignment"], 0.10, 0.15, 3.0, baseline.delay_years, baseline.market_window_open, _ttm_formula(2), "Параметры практического задания"])
        ttm.append([SCENARIO_LABELS["CSV"], baseline.discount_rate, baseline.price_erosion_rate, baseline.market_horizon, baseline.delay_years, baseline.market_window_open, _ttm_formula(3), "Параметры из исходного CSV"])
        ttm.append([SCENARIO_LABELS["Stress Baseline"], 0.10, 0.15, 3.0, stress_baseline_delay, baseline.market_window_open, _ttm_formula(4), "Задержка из базового столбца сравнительного CSV"])
        ttm.append([SCENARIO_LABELS["Stress"], 0.10, 0.15, 3.0, stress.delay_years, stress.market_window_open, _ttm_formula(5), "Задержка из стрессового CSV"])
        for row in (2, 3, 4, 5):
            LOGGER.info("EXCEL_FORMULA_WRITTEN metric=TTM row=%s", row)
            for col in (2, 3, 7):
                ttm.cell(row, col).number_format = "0.0000%"
        _style_sheet(ttm, {"A": 18, "B": 18, "C": 18, "D": 18, "E": 14, "F": 16, "G": 18, "H": 36})

        stress_sheet = workbook.create_sheet(SHEET_NAMES["stress"])
        stress_sheet.append(["Метрика", "Базовый сценарий", "Стрессовый сценарий", "Абсолютное изменение", "Относительное изменение", "Интерпретация"])
        baseline_cost = f"({b['Raw_Material_Cost']}+{b['Energy_kWh']}*{b['Energy_Cost_Rate']}+{b['OPEX_Overhead']}+{b['Equipment_CAPEX']}/{b['Amortization_Years_Economic']}/12)"
        stress_cost = f"({s['Raw_Material_Cost']}+{s['Energy_kWh']}*{s['Energy_Cost_Rate']}+{s['OPEX_Overhead']}+{s['Equipment_CAPEX']}/{s['Amortization_Years_Economic']}/12)"
        stress_rows = [
            (localize_metric("FPY"), f"={b['First_Pass_Good_Units']}/{b['Actual_Output']}", f"={s['First_Pass_Good_Units']}/{s['Actual_Output']}", "Снижается доля годной продукции первого прохода"),
            (localize_metric("OEE"), f"=({b['Actual_Hours']}/{b['Planned_Hours']})*(({b['Planned_Hours']}/{b['Target_Output']})*{b['Actual_Output']}/{b['Actual_Hours']})*({b['First_Pass_Good_Units']}/{b['Actual_Output']})", f"=({s['Actual_Hours']}/{s['Planned_Hours']})*(({s['Planned_Hours']}/{s['Target_Output']})*{s['Actual_Output']}/{s['Actual_Hours']})*({s['First_Pass_Good_Units']}/{s['Actual_Output']})", "Снижение качества уменьшает общую эффективность оборудования"),
            (localize_metric("TEEP"), f"=B3*('{SHEET_NAMES['source']}'!$D$3/'{SHEET_NAMES['source']}'!$D$2)", f"=C3*('{SHEET_NAMES['source']}'!$D$21/'{SHEET_NAMES['source']}'!$D$20)", "Использование календарного времени не изменяется"),
            (localize_metric("Energy Cost"), f"={b['Energy_kWh']}*{b['Energy_Cost_Rate']}", f"={s['Energy_kWh']}*{s['Energy_Cost_Rate']}", "Трёхкратный тариф увеличивает затраты на электроэнергию"),
            (localize_metric("Economic Intensity"), f"={baseline_cost}/{b['Final_Good_Units']}", f"={stress_cost}/{s['Final_Good_Units']}", "Рост затрат повышает стоимость ресурсов на литр годной продукции"),
            (localize_metric("CPU"), "=B6", "=C6", "Численно совпадает с экономической ресурсоёмкостью в текущей области расчёта"),
            (localize_metric("TTM Penalty"), f"='{SHEET_NAMES['ttm']}'!G4", f"='{SHEET_NAMES['ttm']}'!G5", "Задержка создаёт мультипликативный экономический штраф"),
        ]
        for metric_name, baseline_formula, stress_formula, interpretation in stress_rows:
            row_number = stress_sheet.max_row + 1
            stress_sheet.append([metric_name, baseline_formula, stress_formula, f"=C{row_number}-B{row_number}", f'=IF(B{row_number}=0,NA(),D{row_number}/ABS(B{row_number}))', interpretation])
            LOGGER.info("EXCEL_FORMULA_WRITTEN stress_metric=%s", metric_name)
        _style_sheet(stress_sheet, {"A": 26, "B": 20, "C": 20, "D": 20, "E": 20, "F": 62})

        harness = workbook.create_sheet(SHEET_NAMES["harness"])
        harness.append(["Метрика", "Значение Excel", "Значение Python", "Абсолютное отклонение", "Относительное отклонение", "Допуск", "Статус", "Вероятная причина", "Временная метка"])
        _style_sheet(harness, {"A": 28, "B": 18, "C": 18, "D": 18, "E": 18, "F": 14, "G": 18, "H": 54, "I": 26})

        opex = workbook.create_sheet(SHEET_NAMES["opex"])
        opex.append(["Категория затрат", "Доля"])
        for category, share in (
            ("Сырье и прекурсоры", 0.25),
            ("Утилиты", 0.20),
            ("Персонал", 0.30),
            ("ТОиР", 0.15),
            ("Комплаенс, качество и метрология", 0.10),
        ):
            opex.append([category, share])
        chart = BarChart()
        chart.title = "Структура OPEX в высокотехнологичном производстве"
        chart.y_axis.title = "Доля"
        chart.x_axis.title = "Категория"
        chart.add_data(Reference(opex, min_col=2, min_row=1, max_row=6), titles_from_data=True)
        chart.set_categories(Reference(opex, min_col=1, min_row=2, max_row=6))
        opex.add_chart(chart, "D2")
        if opex_chart_path and Path(opex_chart_path).is_file():
            opex.add_image(Image(str(opex_chart_path)), "D18")
        _style_sheet(opex, {"A": 40, "B": 16})

        assumptions_sheet = workbook.create_sheet(SHEET_NAMES["assumptions"])
        assumption_headers = ["ID", "Assumption", "Reason", "Impact", "Source", "Status"]
        assumptions_sheet.append(["Код", "Допущение", "Обоснование", "Влияние", "Источник", "Статус"])
        default_assumptions = [
            {"ID": "A01", "Assumption": "Расчётный период — один месяц", "Reason": "Calendar_Hours = 720 = 30 × 24", "Impact": "Месячная амортизация", "Source": "CSV", "Status": "Принято"},
            {"ID": "A02", "Assumption": "Нормативное время цикла = плановое время / плановый выпуск", "Reason": "Отдельное поле отсутствует", "Impact": "Производительность и OEE", "Source": "Лекция + CSV", "Status": "Принято"},
            {"ID": "A03", "Assumption": "Знаменатель CPU — итоговый годный объём", "Reason": "Это объём после итогового контроля", "Impact": "CPU и ресурсоёмкость", "Source": "Лекция + задание", "Status": "Принято"},
        ]
        for item in assumptions or default_assumptions:
            assumptions_sheet.append([item.get(key, "") for key in assumption_headers])
        _style_sheet(assumptions_sheet, {"A": 10, "B": 52, "C": 48, "D": 30, "E": 24, "F": 16})

        issues_sheet = workbook.create_sheet(SHEET_NAMES["issues"])
        issue_headers = ["Issue", "Parameter", "Severity", "Effect", "Action"]
        issues_sheet.append(["Проблема", "Параметр", "Критичность", "Влияние", "Рекомендуемое действие"])
        default_issues = [
            {"Issue": "Отсутствует масса сырья", "Parameter": "Raw_Material_Mass_kg", "Severity": "ВЫСОКАЯ", "Effect": "Массовая ресурсоёмкость не рассчитывается", "Action": "Добавить массу потреблённого сырья в килограммах"},
            {"Issue": "Параметры TTM расходятся", "Parameter": "Discount_Rate; Price_Erosion_Rate", "Severity": "СРЕДНЯЯ", "Effect": "Получаются разные штрафы TTM", "Action": "Показывать сценарии задания и CSV отдельно"},
            {"Issue": "Задержка пилотного сценария отличается от базового comparison CSV", "Parameter": "Delay_Years", "Severity": "СРЕДНЯЯ", "Effect": "Пилотный сценарий не является сценарием без задержки", "Action": "Использовать comparison CSV для проверки стресс-сценария"},
        ]
        for item in data_issues or default_issues:
            issues_sheet.append([item.get(key, "") for key in issue_headers])
        _style_sheet(issues_sheet, {"A": 38, "B": 38, "C": 14, "D": 48, "E": 56})

        workbook.save(path)
        return path

    def write_harness(self, path: Path, entries: Iterable[Any]) -> None:
        """Write comparison rows, then let Excel recalculate the book again."""
        path = Path(path)
        workbook = load_workbook(path, data_only=False)
        sheet = workbook[SHEET_NAMES["harness"]]
        if sheet.max_row > 1:
            sheet.delete_rows(2, sheet.max_row - 1)
        for entry in entries:
            sheet.append(
                [
                    localize_metric(entry.metric),
                    entry.excel_value,
                    entry.python_value,
                    entry.absolute_delta,
                    entry.relative_delta,
                    entry.tolerance,
                    localize_status(entry.status),
                    localize_cause(entry.probable_cause),
                    entry.timestamp,
                ]
            )
        workbook.save(path)


@dataclass(frozen=True, slots=True)
class RecalculationResult:
    """Outcome and auditable stage markers for Excel automation."""

    success: bool
    events: tuple[str, ...]
    error: str | None = None


class ExcelRecalculator:
    """Use installed Microsoft Excel as the formula engine."""

    def __init__(self, excel_factory: Callable[[], Any] | None = None) -> None:
        self._excel_factory = excel_factory

    def recalculate(self, path: Path) -> RecalculationResult:
        """Force a full Excel calculation, save caches, and close cleanly."""
        path = Path(path).resolve()
        events: list[str] = ["EXCEL_RECALC_STARTED"]
        LOGGER.info("EXCEL_RECALC_STARTED path=%s", path)
        if self._excel_factory is None:
            return self._recalculate_in_worker(path, events)

        return self._recalculate_direct(path, events)

    def _recalculate_in_worker(
        self, path: Path, events: list[str]
    ) -> RecalculationResult:
        """Isolate Excel's COM lifetime from subsequent workbook reads."""
        worker_path = Path(__file__).with_name("excel_worker.py")
        try:
            completed = subprocess.run(
                [sys.executable, str(worker_path), str(path)],
                capture_output=True,
                text=True,
                timeout=120,
                check=False,
            )
        except (OSError, subprocess.SubprocessError) as exc:
            LOGGER.exception("Excel worker could not run for %s", path)
            return RecalculationResult(False, tuple(events), str(exc))
        if completed.returncode != 0:
            error = completed.stderr.strip() or completed.stdout.strip()
            LOGGER.error("Excel worker failed for %s: %s", path, error)
            return RecalculationResult(False, tuple(events), error)
        events.append("EXCEL_RECALC_COMPLETED")
        LOGGER.info("EXCEL_RECALC_COMPLETED path=%s", path)
        return RecalculationResult(True, tuple(events))

    def _recalculate_direct(
        self, path: Path, events: list[str]
    ) -> RecalculationResult:
        """Run through an injected COM-compatible object for unit testing."""
        excel = None
        workbook = None
        try:
            excel = self._excel_factory()
            excel.Visible = False
            excel.DisplayAlerts = False
            workbook = excel.Workbooks.Open(str(path))
            excel.CalculateFullRebuild()
            workbook.Save()
            workbook.Close(SaveChanges=True)
            workbook = None
            events.append("EXCEL_RECALC_COMPLETED")
            LOGGER.info("EXCEL_RECALC_COMPLETED path=%s", path)
            return RecalculationResult(True, tuple(events))
        except Exception as exc:  # external COM boundary; error is returned and logged
            LOGGER.exception("Excel recalculation failed for %s", path)
            return RecalculationResult(False, tuple(events), str(exc))
        finally:
            if workbook is not None:
                try:
                    workbook.Close(SaveChanges=False)
                except Exception:
                    LOGGER.exception("Failed to close Excel workbook after error")
            if excel is not None:
                try:
                    excel.Quit()
                except Exception:
                    LOGGER.exception("Failed to quit Excel after error")


class ExcelGroundTruthReader:
    """Read only values cached by the spreadsheet calculation engine."""

    def read(self, path: Path) -> dict[str, float | None]:
        """Return the ground-truth metric mapping from cached cells."""
        workbook = load_workbook(Path(path), data_only=True, read_only=True)
        values: dict[str, float | None] = {}
        try:
            for metric_name, (sheet_name, cell_address) in GROUND_TRUTH_CELLS.items():
                value = workbook[sheet_name][cell_address].value
                values[metric_name] = float(value) if isinstance(value, (int, float)) else None
                LOGGER.info("EXCEL_VALUE_READ metric=%s value=%s", metric_name, value)
        finally:
            workbook.close()
        return values
