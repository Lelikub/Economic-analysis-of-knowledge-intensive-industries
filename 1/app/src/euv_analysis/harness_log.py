"""Structured implementation-problem journal used by the deliverables."""

from __future__ import annotations

import json
from dataclasses import dataclass
from pathlib import Path
from typing import Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


HARNESS_LOG_HEADERS = (
    "№",
    "Этап",
    "Запрос к ИИ / задача",
    "Сгенерированный или изменённый фрагмент кода",
    "Обнаруженная ошибка / галлюцинация",
    "Причина",
    "Как исправлено",
    "Результат повторной проверки",
)

HARNESS_LOG_WIDTHS = (6.0, 24.0, 32.0, 40.0, 40.0, 35.0, 42.0, 38.0)


@dataclass(frozen=True, slots=True)
class HarnessLogEntry:
    number: int
    stage: str
    task: str
    code_fragment: str
    issue: str
    cause: str
    resolution: str
    recheck_result: str


class HarnessLogDataError(ValueError):
    """Raised when the structured Harness Log source violates its schema."""


class HarnessLogLoader:
    """Load and validate the implementation journal from UTF-8 JSON."""

    _fields = frozenset(HarnessLogEntry.__dataclass_fields__)
    _text_fields = _fields - {"number"}

    def load(self, path: Path) -> tuple[HarnessLogEntry, ...]:
        try:
            payload = json.loads(Path(path).read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError) as exc:
            raise HarnessLogDataError(f"Cannot load Harness Log: {exc}") from exc

        if not isinstance(payload, list):
            raise HarnessLogDataError("Harness Log must be a JSON array")

        entries: list[HarnessLogEntry] = []
        for index, item in enumerate(payload, start=1):
            if not isinstance(item, dict):
                raise HarnessLogDataError(f"Row {index} must be an object")

            keys = set(item)
            missing = self._fields - keys
            unknown = keys - self._fields
            if missing:
                raise HarnessLogDataError(
                    f"Row {index} has missing fields: {', '.join(sorted(missing))}"
                )
            if unknown:
                raise HarnessLogDataError(
                    f"Row {index} has unknown fields: {', '.join(sorted(unknown))}"
                )

            number = item["number"]
            if type(number) is not int or number <= 0:
                raise HarnessLogDataError(f"Row {index} number must be a positive integer")
            for field in self._text_fields:
                value = item[field]
                if not isinstance(value, str) or not value.strip():
                    raise HarnessLogDataError(
                        f"Row {index} field {field!r} must be a non-empty string"
                    )

            entries.append(HarnessLogEntry(**item))

        actual_numbers = [entry.number for entry in entries]
        expected_numbers = list(range(1, len(entries) + 1))
        if actual_numbers != expected_numbers:
            raise HarnessLogDataError(
                "Harness Log number sequence must be consecutive from 1"
            )

        return tuple(entries)


class HarnessLogWorkbookBuilder:
    """Create the approved eight-column Harness Log workbook or sheet."""

    def build(self, path: Path, entries: Sequence[HarnessLogEntry]) -> Path:
        destination = Path(path)
        destination.parent.mkdir(parents=True, exist_ok=True)

        workbook = Workbook()
        workbook.remove(workbook.active)
        self.add_sheet(workbook, entries)
        workbook.save(destination)
        return destination

    def add_sheet(
        self,
        workbook: Workbook,
        entries: Sequence[HarnessLogEntry],
        *,
        title: str = "Harness Log",
    ):
        sheet = workbook.create_sheet(title)
        sheet.append(HARNESS_LOG_HEADERS)
        for entry in entries:
            sheet.append(
                (
                    entry.number,
                    entry.stage,
                    entry.task,
                    entry.code_fragment,
                    entry.issue,
                    entry.cause,
                    entry.resolution,
                    entry.recheck_result,
                )
            )

        header_fill = PatternFill(fill_type="solid", fgColor="1F4E78")
        header_font = Font(color="FFFFFF", bold=True)
        wrapped_top = Alignment(wrap_text=True, vertical="top")
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
        for row in sheet.iter_rows():
            for cell in row:
                cell.alignment = wrapped_top

        for column, width in zip("ABCDEFGH", HARNESS_LOG_WIDTHS, strict=True):
            sheet.column_dimensions[column].width = width
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        return sheet
