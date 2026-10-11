"""Один воспроизводимый запуск всех десяти этапов задания."""

import csv
import hashlib
import json
import logging
import math
import platform
import shutil
import subprocess
import sys
import time
import traceback
import zipfile
from datetime import datetime, timezone
from decimal import Decimal
from pathlib import Path

import numpy as np
from openpyxl import load_workbook
from docx import Document

from . import __version__
from .analytics import (annual_loss, frechet_bounds, lognormal_mean, lognormal_median,
                        parallel_reliability, series_reliability, tco_annual)
from .artifacts import (build_figures, build_ground_truth, build_harness, build_report,
                        build_scenario_book, read_recalculated_ground_truth)
from .historical import HistoricalIncidentAnalyzer
from .loaders import InputDataLoader
from .simulation import simulate_copula, simulate_independent, simulate_pair, summarize_simulation
from .validators import DataValidator


MAIN = ("Node_1_PGM", "Node_2_UltraPure_HF_Main", "Node_3_Monomer",
        "Node_4_Liquid_Nitrogen", "Node_5_Membrane_Reactor")


def write_json(path, data):
    Path(path).write_text(json.dumps(data, ensure_ascii=False, indent=2, default=str), encoding="utf-8")


def sha256(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


class SupplyChainRiskAnalyzer:
    """Uses normalized master data for topology, cost and stochastic analysis."""
    def __init__(self, data, n=10000, seed=314159):
        self.data = data
        self.n = n
        self.seed = seed
        self.r = [data.nodes[k].reliability for k in MAIN]
        self.backup = data.nodes["Node_2_UltraPure_HF_Backup"].reliability
        self.mu = data.downtime.lognorm_mu
        self.sigma = data.downtime.lognorm_sigma
        self.c_day = data.financial.c_day
        self.c_scrap = data.financial.c_scrap
        self.impact_mean = self.c_day * lognormal_mean(self.mu, self.sigma) + self.c_scrap
        self.tco = tco_annual(data.redundancy.capex_qual, data.redundancy.r_amor,
                              data.redundancy.opex_audit, data.redundancy.delta_c_proc)

    def analytic_independent(self, dual=False, rho=0, stress=False):
        r = list(self.r)
        if stress: r[0] = .8  # Explicit assignment in task PDF, p. 4.
        if dual: r[1] = parallel_reliability(r[1], self.backup, rho)
        r_sys = series_reliability(r)
        p_fail = 1-r_sys
        return {"r_sys": r_sys, "p_fail": p_fail,
                "el": annual_loss(p_fail, self.c_day, self.c_scrap,
                                  lognormal_mean(self.mu, self.sigma)),
                "el_med": annual_loss(p_fail, self.c_day, self.c_scrap,
                                      lognormal_median(self.mu))}

    def simulate(self, kind, seed, duration="lognormal"):
        d = self.data.downtime
        common = (self.n, seed, self.mu, self.sigma, self.c_day, self.c_scrap)
        if kind == "S0":
            return simulate_independent(self.r, *common, duration=duration,
                                        pareto_alpha=d.pareto_alpha, pareto_xm=d.pareto_xm_days)
        if kind == "S0_fixed":
            return simulate_independent(self.r, *common, fixed=True)
        if kind == "S1":
            return simulate_independent(self.r, *common, backup=self.backup)
        if kind == "S2":
            return simulate_pair(self.r, self.backup, .2, *common)
        names = self.data.base_matrix.names
        reliability = np.array([(.8 if kind in ("S3", "S4") and name == "Node_1_PGM"
                                 else self.data.nodes[name].reliability) for name in names])
        matrix = self.data.stress_matrix.values if kind in ("S3", "S4") else self.data.base_matrix.values
        return simulate_copula(reliability, matrix, *common, dual=kind in ("S2_full", "S4"))


class PipelineOrchestrator:
    def __init__(self, data_dir, template, output, n=10000, seed=314159):
        self.data_dir = Path(data_dir).resolve()
        self.template = Path(template).resolve()
        self.output = Path(output).resolve()
        self.n, self.seed = n, seed
        self.output.mkdir(parents=True, exist_ok=True)
        (self.output / "logs").mkdir(exist_ok=True)
        (self.output / "data").mkdir(exist_ok=True)
        (self.output / "figures").mkdir(exist_ok=True)
        self.issues = []
        self.stages = {str(i): "SKIPPED" for i in range(1, 11)}
        self.logger = logging.getLogger("catalyst_risk_pipeline")
        self.logger.setLevel(logging.INFO)
        self.logger.handlers.clear()
        handler = logging.FileHandler(self.output / "logs/pipeline.log", encoding="utf-8")
        handler.setFormatter(logging.Formatter("%(asctime)s UTC %(levelname)s %(message)s"))
        handler.formatter.converter = time.gmtime
        self.logger.addHandler(handler)
        self.input_records = []

    def stage(self, number, status, message):
        self.stages[str(number)] = status
        self.logger.info("STAGE=%02d STATUS=%s %s", number, status, message)

    def issue(self, kind, field, description, severity="WARNING", status="UNRESOLVED", **extra):
        row = {"type": kind, "field": field, "description": description,
               "severity": severity, "status": status, **extra}
        self.issues.append(row)
        self.logger.warning("ISSUE %s %s %s", kind, field, description)

    def run(self):
        started = datetime.now(timezone.utc).isoformat()
        self.logger.info("START version=%s Python=%s OS=%s seed=%d N=%d", __version__, sys.version,
                         platform.platform(), self.seed, self.n)
        results, samples, reconciliation, historical, quality = [], {}, [], {}, []
        excel_status = "NOT_ATTEMPTED"
        tests = []
        try:
            data = InputDataLoader().load(self.data_dir)
            self.input_records = list(data.source_records)
            for path in (self.template, self.data_dir / "Задание_к_практическому_занятию_3_ЭА.pdf",
                         self.template.parent.parent / "lection" / "Лекция 5 ЭА слайды.pdf"):
                if path.is_file():
                    self.input_records.append({"file":path.name, "path":str(path),
                                               "size":path.stat().st_size, "sha256":sha256(path)})
            quality = DataValidator().validate(data)
            self.issues.extend({**item, "stage": "1"} for item in quality)
            for source in self.input_records:
                self.logger.info("INPUT %s", json.dumps(source, ensure_ascii=False, default=str))
            normalized = {
                "nodes": {k: vars(v) for k,v in data.nodes.items()},
                "financial": vars(data.financial), "downtime": vars(data.downtime),
                "redundancy": vars(data.redundancy),
                "matrix_kind": "target_binary_pearson",
                "matrix_nodes": data.base_matrix.names,
                "base_matrix": data.base_matrix.values.tolist(),
                "stress_matrix": data.stress_matrix.values.tolist(),
                "source_records": self.input_records,
            }
            write_json(self.output / "data/normalized_inputs.json", normalized)
            blockers = [x for x in quality if x["severity"] == "BLOCKER"]
            self.stage(1, "PASS" if not blockers else "FAIL", f"nodes={len(data.nodes)} incidents={len(data.incidents)} issues={len(quality)}")
            if blockers:
                raise ValueError(f"Input blockers: {len(blockers)}")
            analyzer = SupplyChainRiskAnalyzer(data, self.n, self.seed)
            ground = self.output / "Ground_Truth.xlsx"
            build_ground_truth(ground, data)
            excel_status, excel_values, excel_detail = read_recalculated_ground_truth(ground)
            if excel_status != "VERIFIED":
                self.issue("DOCUMENT_ERROR", "Ground_Truth.xlsx", excel_detail,
                           status="UNRESOLVED", stage="2")
            self.stage(2, "PASS" if excel_status == "VERIFIED" else "WARNING", excel_detail)
            self.stage(3, "PASS", "Typed model and architecture documented")
            # Independent decimal route, separate from production analytics functions.
            decimal_r = [Decimal(str(data.nodes[k].reliability)) for k in MAIN]
            decimal_r_sys = math.prod(decimal_r)
            decimal_mu = Decimal(str(data.downtime.lognorm_mu))
            decimal_sigma = Decimal(str(data.downtime.lognorm_sigma))
            decimal_mean = (decimal_mu + decimal_sigma**2 / 2).exp()
            decimal_el = (1-decimal_r_sys) * (Decimal(str(data.financial.c_day))*decimal_mean + Decimal(str(data.financial.c_scrap)))
            baseline = analyzer.analytic_independent()
            if abs(float(decimal_r_sys)-baseline["r_sys"]) > 1e-12 or abs(float(decimal_el)-baseline["el"]) > 1e-10:
                self.issue("FORMULA_MISMATCH", "baseline", "Decimal and float analytical paths disagree", "BLOCKER", stage="4")
            specs = [
                ("S0", "Независимый моносинг", "independent"),
                ("S1", "Независимый HF резерв", "independent"),
                ("S2", "HF резерв, парная бинарная корреляция 0.20", "pair_binary"),
                ("S0_full", "Моносинг, полная базовая матрица", "Gaussian_NORTA"),
                ("S2_full", "HF резерв, полная базовая матрица", "Gaussian_NORTA"),
                ("S3", "Стрессовый моносинг, полная матрица", "Gaussian_NORTA"),
                ("S4", "Стрессовый HF резерв, полная матрица", "Gaussian_NORTA"),
                ("S5", "Независимый моносинг, Парето длительности", "independent_pareto"),
            ]
            sample_path = self.output / "data/simulation_results.csv"
            with sample_path.open("w", encoding="utf-8-sig", newline="") as file:
                writer = csv.writer(file)
                writer.writerow(["scenario_id", "trial_id", "system_ok", "downtime_days",
                                 "loss_usd_millions", "seed", "duration_model", "n"])
                scenario_seeds = {"S0":self.seed, "S1":self.seed+1, "S2":self.seed+2,
                                  "S0_full":self.seed+3, "S2_full":self.seed+3,
                                  "S3":self.seed+4, "S4":self.seed+4, "S5":self.seed+6}
                for ix, (sid, desc, method) in enumerate(specs):
                    self.logger.info("SCENARIO_START %s method=%s seed=%d", sid, method, scenario_seeds[sid])
                    try:
                        scenario_seed = scenario_seeds[sid]
                        sample = analyzer.simulate("S0" if sid == "S5" else sid,
                            scenario_seed, duration="pareto" if sid == "S5" else "lognormal")
                        samples[sid] = sample
                        sm = summarize_simulation(sample)
                        analytic = None
                        if sid in ("S0", "S5"):
                            analytic = analyzer.analytic_independent()
                        elif sid == "S1":
                            analytic = analyzer.analytic_independent(dual=True)
                        elif sid == "S2":
                            analytic = analyzer.analytic_independent(dual=True, rho=.2)
                        row = {"id": sid, "description": desc, "method": method, "status": "PASS",
                               "n": self.n, "seed": scenario_seed,
                               "r_sys": analytic["r_sys"] if analytic else sm["r_sys_mc"],
                               "p_fail": analytic["p_fail"] if analytic else sm["p_fail_mc"],
                               "el": analytic["el"] if analytic and sid != "S5" else sm["el_mc"],
                               "el_mc": sm["el_mc"], "r_sys_mc": sm["r_sys_mc"],
                               "r_se": sm["r_se"], "el_se": sm["el_se"],
                               "tco": analyzer.tco if sid in ("S1", "S2", "S2_full", "S4") else 0,
                               "var95": sm["var95"], "cvar95": sm["cvar95"],
                               "var99": sm["var99"], "cvar99": sm["cvar99"]}
                        if sid == "S5":
                            mean_pareto = data.downtime.pareto_alpha * data.downtime.pareto_xm_days/(data.downtime.pareto_alpha-1)
                            row["el_analytic_pareto"] = annual_loss(analytic["p_fail"], analyzer.c_day, analyzer.c_scrap, mean_pareto)
                            row["el"] = row["el_analytic_pareto"]
                            row["status"] = "WARNING"  # Infinite theoretical variance of the Pareto mean.
                        if "diagnostics" in sample:
                            row["copula_diagnostics"] = sample["diagnostics"]
                            binary = sample["states"].astype(float)
                            empirical = np.corrcoef(binary, rowvar=False)
                            target = (data.stress_matrix if sid in ("S3", "S4") else data.base_matrix).values
                            row["max_binary_corr_error"] = float(np.max(np.abs(empirical-target)))
                        results.append(row)
                        for trial in range(self.n):
                            writer.writerow([sid, trial, int(sample["ok"][trial]),
                                             float(sample["downtime"][trial]), float(sample["loss"][trial]),
                                             scenario_seed, "pareto" if sid == "S5" else "lognormal", self.n])
                        self.logger.info("SCENARIO_END %s %s", sid, json.dumps({k:v for k,v in row.items() if k not in ("description",)}, default=str))
                    except Exception as exc:
                        self.issue("MATH_INVALID", sid, f"Scenario blocked: {type(exc).__name__}: {exc}",
                                   "BLOCKER", stage="7/8")
                        results.append({"id": sid, "description": desc, "method": method,
                                        "status": "BLOCKED", "r_sys": None, "p_fail": None, "el": None,
                                        "tco": analyzer.tco if sid in ("S1", "S2", "S2_full", "S4") else 0})
            self.stage(4, "PASS", f"simulation file={sample_path} trials={sum(len(x['loss']) for x in samples.values())}")
            fixed_sample = analyzer.simulate("S0_fixed", self.seed+100)
            fixed_summary = summarize_simulation(fixed_sample)
            baseline_fixed = baseline["el_med"]
            self.logger.info("S0_FIXED analytic=%s MC=%s", baseline_fixed, fixed_summary["el_mc"])
            test_process = subprocess.run([sys.executable, "-m", "pytest", "tests", "-q"],
                cwd=Path(__file__).resolve().parents[1], capture_output=True, text=True, encoding="utf-8", errors="replace")
            (self.output / "logs/tests.log").write_text(test_process.stdout + test_process.stderr, encoding="utf-8")
            passed = test_process.returncode == 0
            tests.append(["PYTEST", "Core mathematical and IO invariants", "tests/", "exit=0",
                          f"exit={test_process.returncode}; {test_process.stdout.strip()}",
                          "PASS" if passed else "FAIL", "logs/tests.log"])
            self.stage(5, "PASS" if passed else "FAIL", test_process.stdout.strip())
            if not passed:
                self.issue("CODE_ERROR", "pytest", test_process.stdout + test_process.stderr, "BLOCKER", stage="5")
            byid = {r["id"]: r for r in results}
            for sid, refid in (("S1", "S0"), ("S2", "S0"), ("S2_full", "S0_full"), ("S4", "S3")):
                row, ref = byid[sid], byid[refid]
                if row.get("el") is not None and ref.get("el") is not None:
                    row["avoided"] = ref["el"]-row["el"]
                    row["net_effect"] = row["avoided"]-analyzer.tco
                    row["roi"] = row["net_effect"]/analyzer.tco
                    if sid in ("S2_full", "S4") and sid in samples and refid in samples:
                        paired = samples[refid]["loss"] - samples[sid]["loss"]
                        se = float(np.std(paired, ddof=1) / np.sqrt(len(paired)))
                        row["avoided_se"] = se
                        row["avoided_ci95_low"] = row["avoided"] - 1.96*se
                        row["avoided_ci95_high"] = row["avoided"] + 1.96*se
                        row["roi_ci95_low"] = (row["avoided_ci95_low"]-analyzer.tco)/analyzer.tco
                        row["roi_ci95_high"] = (row["avoided_ci95_high"]-analyzer.tco)/analyzer.tco
            self.stage(6, "PASS", f"TCO={analyzer.tco:.4f}; S1 ROI={byid['S1'].get('roi')}")
            self.stage(7, "PASS" if byid["S0_full"]["status"] != "BLOCKED" and byid["S2_full"]["status"] != "BLOCKED" else "BLOCKED",
                       f"S2 pair/full={byid['S2']['status']}/{byid['S2_full']['status']}; matched baseline={byid['S0_full']['status']}")
            self.stage(8, "PASS" if byid["S3"]["status"] != "BLOCKED" and byid["S4"]["status"] != "BLOCKED" else "BLOCKED",
                       f"S3/S4={byid['S3']['status']}/{byid['S4']['status']}")
            hist_analyzer = HistoricalIncidentAnalyzer().set_prior(analyzer.mu, analyzer.sigma,
                        data.downtime.pareto_alpha, data.downtime.pareto_xm_days)
            historical = hist_analyzer.analyze(data.incidents, analyzer.c_day, analyzer.c_scrap, self.seed)
            tail_sensitivity = []
            for duration in ("lognormal", "pareto"):
                for n_sens in (2000, self.n):
                    for seed_sens in (self.seed+200, self.seed+201, self.seed+202):
                        sample_sens = simulate_independent(analyzer.r, n_sens, seed_sens,
                            analyzer.mu, analyzer.sigma, analyzer.c_day, analyzer.c_scrap,
                            duration=duration, pareto_alpha=data.downtime.pareto_alpha,
                            pareto_xm=data.downtime.pareto_xm_days)
                        summary_sens = summarize_simulation(sample_sens)
                        tail_sensitivity.append({"duration":duration, "n":n_sens,
                            "seed":seed_sens, "el":summary_sens["el_mc"],
                            "var99":summary_sens["var99"], "cvar99":summary_sens["cvar99"]})
            historical["tail_sensitivity"] = tail_sensitivity
            write_json(self.output / "data/historical_analysis.json", historical)
            write_json(self.output / "data/tail_sensitivity.json", tail_sensitivity)
            if historical["residual_nonzero"]:
                self.issue("STATISTICAL_WARNING", "historical loss residuals",
                           f"{historical['residual_nonzero']} recorded losses differ from C_day*T+C_scrap; mean residual {historical['residual_mean']:.4f} USD m",
                           stage="9", evidence="data/historical_analysis.json", cause="Source loss may include other components; not documented")
            self.stage(9, "PASS", f"incidents={historical['count']} Shapiro p={historical['shapiro_p']:.4f}")
            # Independent analytic and simulation checks, with sampling error treated separately.
            for metric, analytic, mc, se, excel_key in (
                ("R_sys", baseline["r_sys"], byid["S0"]["r_sys_mc"], byid["S0"]["r_se"], "R_sys"),
                ("P_fail", baseline["p_fail"], 1-byid["S0"]["r_sys_mc"], byid["S0"]["r_se"], "P_fail"),
                ("EL_med", baseline["el_med"], fixed_summary["el_mc"], fixed_summary["el_se"], "EL_med"),
                ("EL_lognorm", baseline["el"], byid["S0"]["el_mc"], byid["S0"]["el_se"], "EL_lognorm"),
            ):
                excel = excel_values.get(excel_key)
                tolerance = 4*se + 1e-8
                abs_dev = abs(mc-analytic)
                status = "PASS" if abs_dev <= tolerance else "FAIL"
                if excel_status != "VERIFIED" and status == "PASS": status = "NOT_VERIFIED_EXCEL"
                reconciliation.append({"metric":metric, "scenario":"S0", "excel":excel,
                    "analytic":analytic, "mc":mc, "abs_dev":abs_dev,
                    "rel_dev":abs_dev/abs(analytic) if analytic else "N/A",
                    "tolerance":f"4×SE={tolerance:.6g}; SE={se:.6g}", "status":status})
            for metric, analytic, excel_key in (("R_sys", float(decimal_r_sys), "R_sys"),
                                                 ("EL_lognorm", float(decimal_el), "EL_lognorm")):
                value = baseline["r_sys"] if metric == "R_sys" else baseline["el"]
                reconciliation.append({"metric":metric+" independent Decimal", "scenario":"S0",
                    "excel":excel_values.get(excel_key), "analytic":value, "mc":analytic,
                    "abs_dev":abs(value-analytic), "rel_dev":abs(value-analytic)/abs(value),
                    "tolerance":"1e-10 deterministic", "status":"PASS" if abs(value-analytic)<1e-10 else "FAIL"})
            workbook = self.output / "Scenario_Results.xlsx"
            build_scenario_book(workbook, results, quality, historical, tail_sensitivity)
            figures = build_figures(self.output / "figures", results, samples, historical, data)
            metadata = {"timestamp_utc":started, "version":__version__, "python":sys.version.split()[0],
                        "seed":self.seed, "N":self.n, "excel_status":excel_status,
                        "sources":self.input_records, "stage_statuses":self.stages}
            report = self.output / "Отчет_Практическая_работа_3_Вариант_Б.docx"
            build_report(report, self.template, results, historical, self.issues, quality,
                         excel_status, figures, metadata)
            self._write_memo(results, historical, excel_status)
            self.stage(10, "PASS", f"report={report.name} scenarios={len(results)} figures={len(figures)}")
            self._write_docs(data, results, historical)
        except Exception as exc:
            self.logger.error("PIPELINE_ERROR %s\n%s", exc, traceback.format_exc())
            self.issue("CODE_ERROR", "pipeline", f"{type(exc).__name__}: {exc}", "BLOCKER",
                       stage="orchestrator", evidence=traceback.format_exc())
        finally:
            metadata = {"timestamp_utc":started, "finished_utc":datetime.now(timezone.utc).isoformat(),
                        "version":__version__, "python":sys.version.split()[0], "seed":self.seed,
                        "N":self.n, "excel_status":excel_status, "stage_statuses":self.stages}
            assumptions = [
                ["A1", "One annual state and at most one incident", "Task prompt §11", "All scenarios", "No repeated-failure frequency", "MODEL"],
                ["A2", "CSV rho targets binary Pearson correlation", "Task prompt §9", "Copula scenarios", "Requires NORTA calibration", "MODEL"],
                ["A3", "Other stages independent in S2 pair", "Task prompt §9", "S2", "Full matrix differs", "MODEL"],
                ["A4", "PGM R=0.80 under stress", "Assignment p.4", "S3/S4", "Changes marginal reliability", "SOURCE"],
            ]
            try:
                build_harness(self.output / "HARNES_LOG.xlsx", self.issues, reconciliation,
                              tests, quality, assumptions, metadata)
            except Exception as exc:
                self.logger.error("HARNESS_SAVE_ERROR %s", exc)
            self._finish_manifest(started)
        return self.output

    def _write_memo(self, results, historical, excel_status):
        byid = {x["id"]:x for x in results}
        s0, s1, s2, s2_full, s3, s4 = (byid[k] for k in ("S0", "S1", "S2", "S2_full", "S3", "S4"))
        recommendation = "Условно рекомендовать" if s1.get("roi", -1) > 0 else "Не рекомендовать"
        lines = ["# Управленческий меморандум", "", f"**Решение:** {recommendation.lower()} HF-резерв при документально подтверждённой независимости поставщиков.", "",
                 f"- S0: R_sys={s0['r_sys']:.5f}; EL={s0['el']:.3f} млн USD/год; VaR99={s0['var99']:.3f}; CVaR99={s0['cvar99']:.3f}.",
                 f"- S1: R_sys={s1['r_sys']:.5f}; предотвращённый ущерб={s1.get('avoided', float('nan')):.3f} млн USD/год; TCO={s1['tco']:.4f}; ROI={s1.get('roi', float('nan')):.3f}.",
                 f"- Доля TCO от бюджета снабжения 120 млн USD: {100*s1['tco']/120:.2f}%; доля предотвращённого ущерба: {100*s1.get('avoided',0)/120:.2f}%. Это сопоставление масштабов, а не разрешение расходовать бюджет.",
                 f"- S2, парная корреляция HF ρ=0.20: R_sys={s2.get('r_sys',float('nan')):.5f}, ROI={s2.get('roi',float('nan')):.3f}; эффект ниже независимого резерва.",
                 f"- S2_full против S0_full, полная базовая матрица при одинаковом seed: предотвращённый ущерб={s2_full.get('avoided',float('nan')):.3f} млн USD/год, ROI={s2_full.get('roi',float('nan')):.3f}; выборочный 95% интервал ROI [{s2_full.get('roi_ci95_low',float('nan')):.3f}; {s2_full.get('roi_ci95_high',float('nan')):.3f}].",
                 f"- S3/S4, стресс R_PGM=0.80 и полная матрица: EL={s3.get('el',float('nan')):.3f}/{s4.get('el',float('nan')):.3f} млн USD/год, ROI резерва={s4.get('roi',float('nan')):.3f}; выборочный 95% интервал ROI [{s4.get('roi_ci95_low',float('nan')):.3f}; {s4.get('roi_ci95_high',float('nan')):.3f}].",
                 f"- Хвост S0: CVaR99−VaR99={s0['cvar99']-s0['var99']:.3f} млн USD; средний EL не характеризует редкие крупные потери.",
                 f"- Excel Ground Truth: {excel_status}; 50 исторических инцидентов не дают частоту годовых отказов.",
                 "- Улучшения методики: оценить процесс повторных отказов, калибровать зависимость по наблюдениям, увеличить N и число seed для CVaR99, провести экспертизу модели простоев."]
        (self.output / "Management_Memo.md").write_text("\n".join(lines), encoding="utf-8")

    def _write_docs(self, data, results, historical):
        source = Path(__file__).resolve().parents[1]
        for name in ("README.md", "ARCHITECTURE.md", "requirements.txt"):
            shutil.copy2(source / name, self.output / name)
        shutil.copytree(source / "catalyst_risk_pipeline", self.output / "catalyst_risk_pipeline", dirs_exist_ok=True,
                        ignore=shutil.ignore_patterns("__pycache__", "*.pyc"))
        shutil.copytree(source / "tests", self.output / "tests", dirs_exist_ok=True,
                        ignore=shutil.ignore_patterns("__pycache__", "*.pyc"))

    def _finish_manifest(self, started):
        self.logger.info("FINALIZING stages=%s issue_count=%d", self.stages, len(self.issues))
        for handler in self.logger.handlers:
            handler.flush()
        files = []
        for path in sorted(self.output.rglob("*")):
            if path.is_file() and path.name not in ("run_manifest.json", "Practical_3_Variant_B_FULL.zip"):
                try:
                    if path.suffix == ".xlsx":
                        wb = load_workbook(path, read_only=True)
                        wb.close()
                    if path.suffix == ".docx":
                        Document(path)
                    files.append({"path":str(path.relative_to(self.output)).replace("\\", "/"),
                                  "bytes":path.stat().st_size, "sha256":sha256(path), "open_check":"PASS"})
                except Exception as exc:
                    files.append({"path":str(path.relative_to(self.output)), "open_check":f"FAIL:{exc}"})
                    self.issue("DOCUMENT_ERROR", str(path), str(exc), "BLOCKER")
        bad = any(x["open_check"] != "PASS" for x in files)
        all_stages = all(x == "PASS" for x in self.stages.values())
        status = "SUCCESS" if all_stages and not bad and not any(x["severity"] == "BLOCKER" for x in self.issues) else "PARTIAL"
        manifest = {"status":status, "started_utc":started,
                    "finished_utc":datetime.now(timezone.utc).isoformat(),
                    "version":__version__, "seed":self.seed, "N":self.n,
                    "stages":self.stages, "sources":self.input_records,
                    "issues":self.issues, "files":files}
        write_json(self.output / "run_manifest.json", manifest)
        archive = self.output / "Practical_3_Variant_B_FULL.zip"
        with zipfile.ZipFile(archive, "w", compression=zipfile.ZIP_DEFLATED) as z:
            for item in files:
                if item["open_check"] == "PASS":
                    z.write(self.output / item["path"], item["path"])
            z.write(self.output / "run_manifest.json", "run_manifest.json")
