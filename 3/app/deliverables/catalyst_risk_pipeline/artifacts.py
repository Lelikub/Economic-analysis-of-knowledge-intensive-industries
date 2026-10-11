"""Книги, графики и отчёт из машинных результатов расчёта."""

import json
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .analytics import frechet_bounds, parallel_reliability


def _style_sheet(ws):
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor="17365D")
        cell.font = Font(color="FFFFFF", bold=True)
    for col in ws.columns:
        letter = col[0].column_letter
        width = min(55, max(14, max(len(str(c.value or "")) for c in col) + 2))
        ws.column_dimensions[letter].width = width


def _table_sheet(wb, name, columns, rows):
    ws = wb.create_sheet(name)
    ws.append(columns)
    for row in rows:
        ws.append([json.dumps(x, ensure_ascii=False) if isinstance(x, (dict, list, tuple)) else x for x in row])
    _style_sheet(ws)
    return ws


def build_ground_truth(path, data):
    wb = Workbook()
    inp = wb.active
    inp.title = "Inputs"
    inp.append(["Parameter", "Value", "Source file / field", "Unit"])
    main = ("Node_1_PGM", "Node_2_UltraPure_HF_Main", "Node_3_Monomer",
            "Node_4_Liquid_Nitrogen", "Node_5_Membrane_Reactor")
    for key in main:
        inp.append([key, data.nodes[key].reliability, f"variant_b_nodes_parameters.csv:{key}:Reliability_R", "probability"])
    values = [
        ("C_day", data.financial.c_day, "variant_b_financial_parameters.csv:C_day", "USD m/day"),
        ("C_scrap", data.financial.c_scrap, "variant_b_financial_parameters.csv:C_scrap", "USD m/incident"),
        ("lognorm_mu", data.downtime.lognorm_mu, "variant_b_financial_parameters.csv:Lognormal_mu", "log days"),
        ("lognorm_sigma", data.downtime.lognorm_sigma, "variant_b_financial_parameters.csv:Lognormal_sigma", "dimensionless"),
        ("HF_backup_R", data.nodes["Node_2_UltraPure_HF_Backup"].reliability, "variant_b_nodes_parameters.csv:HF_Backup:Reliability_R", "probability"),
        ("CAPEX", data.redundancy.capex_qual, "variant_b_redundancy_tco.csv:CAPEX_qual", "USD m"),
        ("r_amor", data.redundancy.r_amor, "variant_b_redundancy_tco.csv:r_amor", "1/year"),
        ("audit", data.redundancy.opex_audit, "variant_b_redundancy_tco.csv:OPEX_audit", "USD m/year"),
        ("procurement", data.redundancy.delta_c_proc, "variant_b_redundancy_tco.csv:Delta_C_proc", "USD m/year"),
    ]
    for item in values:
        inp.append(item)
    # Rows 2:6 nodes, 7:15 financial/downtime/redundancy.
    base = wb.create_sheet("Baseline")
    for row in [
        ["Metric", "Formula", "Unit"], ["R_sys", "=PRODUCT(Inputs!B2:B6)", "probability"],
        ["P_fail", "=1-B2", "probability"],
        ["Impact_med", "=Inputs!B7*Downtime!B2+Inputs!B8", "USD m/incident"],
        ["EL_med", "=B3*B4", "USD m/year"],
        ["EL_lognorm", "=B3*(Inputs!B7*Downtime!B3+Inputs!B8)", "USD m/year"],
    ]: base.append(row)
    down = wb.create_sheet("Downtime")
    for row in [["Metric", "Formula", "Unit"],
                ["Median", "=EXP(Inputs!B9)", "days"],
                ["Mean", "=EXP(Inputs!B9+Inputs!B10^2/2)", "days"]]: down.append(row)
    red = wb.create_sheet("Redundancy")
    for row in [["Metric", "Formula", "Unit"],
                ["R_HF_independent", "=1-(1-Inputs!B3)*(1-Inputs!B11)", "probability"],
                ["R_sys_dual", "=Inputs!B2*B2*PRODUCT(Inputs!B4:B6)", "probability"],
                ["P_fail_dual", "=1-B3", "probability"],
                ["TCO_annual", "=Inputs!B12*Inputs!B13+Inputs!B14+Inputs!B15", "USD m/year"],
                ["Avoided_EL_lognorm", "=(B3-Baseline!B2)*(Inputs!B7*Downtime!B3+Inputs!B8)", "USD m/year"],
                ["Net_effect", "=B6-B5", "USD m/year"],
                ["ROI", "=B7/B5", "ratio"]]: red.append(row)
    fre = wb.create_sheet("Frechet")
    fre.append(["Pair", "R_a", "R_b", "rho_min", "rho_max", "Source"])
    for row in data.frechet_rows:
        a, b = row["Pair_Node_1"], row["Pair_Node_2"]
        r = fre.max_row + 1
        fre.append([f"{a} / {b}", data.nodes[a].reliability, data.nodes[b].reliability,
                    f"=(MAX(0,(1-B{r})+(1-C{r})-1)-(1-B{r})*(1-C{r}))/SQRT(B{r}*(1-B{r})*C{r}*(1-C{r}))",
                    f"=(MIN(1-B{r},1-C{r})-(1-B{r})*(1-C{r}))/SQRT(B{r}*(1-B{r})*C{r}*(1-C{r}))",
                    "variant_b_frechet_bounds.csv; formula lecture 5 pp. 10, 20"])
    cross = wb.create_sheet("Cross_Check")
    for row in [["Metric", "Excel formula link", "Unit"],
                ["R_sys", "=Baseline!B2", "probability"], ["P_fail", "=Baseline!B3", "probability"],
                ["t_med", "=Downtime!B2", "days"], ["E_T", "=Downtime!B3", "days"],
                ["EL_med", "=Baseline!B5", "USD m/year"],
                ["EL_lognorm", "=Baseline!B6", "USD m/year"],
                ["R_sys_dual", "=Redundancy!B3", "probability"],
                ["TCO", "=Redundancy!B5", "USD m/year"]]: cross.append(row)
    for ws in wb:
        _style_sheet(ws)
    wb.calculation.fullCalcOnLoad = True
    wb.save(path)


def read_recalculated_ground_truth(path):
    try:
        import win32com.client
        excel = win32com.client.DispatchEx("Excel.Application")
        excel.Visible = False
        excel.DisplayAlerts = False
        try:
            book = excel.Workbooks.Open(str(Path(path).resolve()))
            excel.CalculateFullRebuild()
            book.Save()
            book.Close(SaveChanges=True)
        finally:
            excel.Quit()
        wb = load_workbook(path, data_only=True, read_only=True)
        keys = ("R_sys", "P_fail", "t_med", "E_T", "EL_med", "EL_lognorm", "R_sys_dual", "TCO")
        values = {key: wb["Cross_Check"].cell(i+2, 2).value for i, key in enumerate(keys)}
        wb.close()
        if not all(isinstance(x, (int, float)) for x in values.values()):
            return "NOT_VERIFIED", {}, "Excel returned empty formula cache"
        return "VERIFIED", values, "Microsoft Excel COM full recalculation"
    except Exception as exc:
        return "NOT_VERIFIED", {}, f"Independent Excel calculation unavailable: {type(exc).__name__}: {exc}"


def build_scenario_book(path, results, data_quality, historical, tail_sensitivity=()):
    wb = Workbook()
    wb.remove(wb.active)
    columns = ["Scenario", "Description", "Method", "R_sys", "P_fail", "EL_mUSD_year",
               "TCO_mUSD_year", "Avoided_mUSD_year", "ROI", "VaR95", "CVaR95", "VaR99", "CVaR99", "Status",
               "Avoided_SE", "Avoided_CI95_low", "Avoided_CI95_high", "ROI_CI95_low", "ROI_CI95_high"]
    rows = [[x["id"], x["description"], x["method"], x.get("r_sys"), x.get("p_fail"),
             x.get("el"), x.get("tco"), x.get("avoided"), x.get("roi"), x.get("var95"),
             x.get("cvar95"), x.get("var99"), x.get("cvar99"), x["status"],
             x.get("avoided_se"), x.get("avoided_ci95_low"), x.get("avoided_ci95_high"),
             x.get("roi_ci95_low"), x.get("roi_ci95_high")] for x in results]
    _table_sheet(wb, "Scenarios", columns, rows)
    _table_sheet(wb, "Data_Quality", ["Type", "Field", "Description", "Severity", "Status"],
                 [[x.get(k) for k in ("type", "field", "description", "severity", "status")] for x in data_quality])
    _table_sheet(wb, "Historical_Residuals", ["Incident_ID", "Observed", "Model", "Residual"],
                 [[x[k] for k in ("id", "observed", "model", "residual")] for x in historical["residual_rows"]])
    _table_sheet(wb, "Tail_Sensitivity", ["Duration", "N", "Seed", "EL", "VaR99", "CVaR99"],
                 [[x[k] for k in ("duration", "n", "seed", "el", "var99", "cvar99")] for x in tail_sensitivity])
    wb.save(path)


def build_harness(path, issues, reconciliation, tests, data_quality, assumptions, metadata):
    wb = Workbook()
    wb.remove(wb.active)
    headers = ["Issue_ID", "UTC", "Stage", "File/key", "Task", "Fragment", "Type", "Description",
               "Evidence", "Cause", "Severity", "Correction", "Retest", "Status", "Log/test"]
    rows = []
    for i, issue in enumerate(issues, 1):
        rows.append([f"ISS-{i:03}", metadata["timestamp_utc"], issue.get("stage", "Input/analysis"),
                     issue.get("field", ""), "Практическая работа 3", issue.get("fragment", ""),
                     issue["type"], issue["description"], issue.get("evidence", issue["description"]),
                     issue.get("cause", ""), issue.get("severity", "WARNING"),
                     issue.get("correction", "Not changed"), issue.get("retest", "Pending"),
                     issue.get("status", "UNRESOLVED"), issue.get("link", "logs/pipeline.log")])
    if not rows:
        rows.append(["ISS-000", metadata["timestamp_utc"], "All", "", "", "", "INFO",
                     "Подтверждённых ошибок не обнаружено", "", "", "INFO", "", "", "PASS", ""])
    _table_sheet(wb, "Issues", headers, rows)
    _table_sheet(wb, "Reconciliation", ["Metric", "Scenario", "Excel", "Analytic Python", "Monte Carlo",
                                        "Abs deviation", "Relative deviation", "CI / tolerance", "Status"],
                 [[x.get(k) for k in ("metric", "scenario", "excel", "analytic", "mc", "abs_dev", "rel_dev", "tolerance", "status")] for x in reconciliation])
    _table_sheet(wb, "Tests", ["ID", "Invariant", "Input", "Expected", "Actual", "Status", "Evidence"], tests)
    _table_sheet(wb, "Data_Quality", ["Problem", "Field", "Severity", "Impact", "Resolution", "Status"],
                 [[x["type"], x["field"], x["severity"], x["description"], x.get("correction", "Review"), x["status"]] for x in data_quality])
    _table_sheet(wb, "Assumptions", ["ID", "Assumption", "Source", "Scope", "Impact", "Status"], assumptions)
    _table_sheet(wb, "Run_Metadata", ["Key", "Value"], [[k, v] for k, v in metadata.items()])
    wb.save(path)


def build_figures(directory, results, samples, historical, data):
    directory = Path(directory)
    directory.mkdir(parents=True, exist_ok=True)
    plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 9, "figure.dpi": 140})
    paths = []
    fig, ax = plt.subplots(figsize=(12, 2.8))
    labels = ["PGM\nSPOF", "HF main\n+ HF reserve", "Monomer", "Liquid N₂", "Reactor"]
    for i, label in enumerate(labels):
        ax.text(i, 0, label, ha="center", va="center", bbox={"boxstyle":"round", "facecolor":"#dceaf5", "edgecolor":"#17365d"})
        if i: ax.annotate("", xy=(i-.27, 0), xytext=(i-.73, 0), arrowprops={"arrowstyle":"->"})
    ax.set_xlim(-.6, 4.6); ax.set_ylim(-.7, .7); ax.axis("off")
    fig.tight_layout(); path = directory / "chain.png"; fig.savefig(path, dpi=220); plt.close(fig); paths.append(path)
    good = [x for x in results if x["status"] in ("PASS", "WARNING") and x.get("el") is not None]
    fig, axes = plt.subplots(1, 2, figsize=(11, 4))
    positions = np.arange(len(good))
    axes[0].bar(positions-.18, [x["r_sys"] for x in good], width=.35, label="R_sys", color="#2c6e91")
    axes[0].bar(positions+.18, [x["p_fail"] for x in good], width=.35, label="P_fail", color="#d68b45")
    axes[0].set_xticks(positions, [x["id"] for x in good]); axes[0].legend()
    axes[0].set_ylabel("Вероятность за год")
    axes[1].bar([x["id"] for x in good], [x["el"] for x in good], color="#a64c3c")
    axes[1].set_ylabel("Ожидаемый ущерб, млн USD/год")
    fig.tight_layout(); path = directory / "scenarios.png"; fig.savefig(path, dpi=220); plt.close(fig); paths.append(path)
    a = data.nodes["Node_2_UltraPure_HF_Main"].reliability
    b = data.nodes["Node_2_UltraPure_HF_Backup"].reliability
    lo, hi = frechet_bounds(a, b)
    x = np.linspace(lo, hi, 200)
    other = np.prod([data.nodes[k].reliability for k in ("Node_1_PGM", "Node_3_Monomer", "Node_4_Liquid_Nitrogen", "Node_5_Membrane_Reactor")])
    impact = data.financial.c_day * np.exp(data.downtime.lognorm_mu + data.downtime.lognorm_sigma**2/2) + data.financial.c_scrap
    base = other * a
    dual = np.array([other * parallel_reliability(a, b, v) for v in x])
    roi = ((dual-base)*impact - data.redundancy.stated_tco) / data.redundancy.stated_tco
    fig, axes = plt.subplots(1, 2, figsize=(11, 4))
    axes[0].plot(x, dual/other); axes[0].set_ylabel("R блока HF")
    axes[1].plot(x, roi); axes[1].axhline(0, color="black", lw=.8); axes[1].set_ylabel("ROI")
    for ax in axes: ax.set_xlabel("Бинарная корреляция Пирсона ρ")
    fig.tight_layout(); path = directory / "correlation.png"; fig.savefig(path, dpi=220); plt.close(fig); paths.append(path)
    fig, axes = plt.subplots(1, 2, figsize=(11, 4))
    for key in ("S0", "S1", "S4"):
        if key in samples:
            v = samples[key]["loss"]
            axes[0].hist(v, bins=40, density=True, histtype="step", label=key)
            axes[1].plot(np.sort(v), np.arange(1, len(v)+1)/len(v), label=key)
    axes[0].set_xlabel("Годовой ущерб, млн USD"); axes[0].set_ylabel("Плотность")
    if "S0" in samples:
        from .analytics import var, cvar
        axes[0].axvline(var(samples["S0"]["loss"], .99), color="#b03a2e", ls="--", label="S0 VaR99")
        axes[0].axvline(cvar(samples["S0"]["loss"], .99), color="#7d3c98", ls=":", label="S0 CVaR99")
    axes[1].set_xlabel("Годовой ущерб, млн USD"); axes[1].set_ylabel("Эмпирическая CDF")
    axes[0].legend(); axes[1].legend(); fig.tight_layout()
    path = directory / "loss_distributions.png"; fig.savefig(path, dpi=220); plt.close(fig); paths.append(path)
    durations = np.array([x.downtime_days for x in data.incidents])
    from scipy import stats
    fig, axes = plt.subplots(1, 3, figsize=(13, 3.8))
    axes[0].hist(durations, bins=12, density=True, alpha=.6)
    grid = np.linspace(max(.01, durations.min()), durations.max(), 200)
    axes[0].plot(grid, stats.lognorm.pdf(grid, historical["sigma_fitted"], scale=np.exp(historical["mu_fitted"])))
    axes[0].set_xlabel("Простой, дни"); axes[0].set_ylabel("Плотность")
    qq = stats.probplot(np.log(durations), dist="norm")
    axes[1].scatter(qq[0][0], qq[0][1], s=18)
    axes[1].plot(qq[0][0], qq[1][1] + qq[1][0]*np.asarray(qq[0][0]))
    axes[1].set_xlabel("Нормальные квантили"); axes[1].set_ylabel("log(простой)")
    grid = np.linspace(durations.min(), durations.max(), 200)
    axes[2].step(np.sort(durations), np.arange(1, len(durations)+1)/len(durations), where="post", label="История")
    axes[2].plot(grid, stats.lognorm.cdf(grid, historical["sigma_fitted"], scale=np.exp(historical["mu_fitted"])), label="Логнорм.")
    pareto = historical["pareto"]
    if "fitted_alpha" in pareto:
        axes[2].plot(grid, stats.pareto.cdf(grid, pareto["fitted_alpha"], scale=pareto["threshold_days"]), label="Парето")
    axes[2].set_xlabel("Простой, дни"); axes[2].set_ylabel("CDF"); axes[2].legend()
    fig.tight_layout(); path = directory / "historical.png"; fig.savefig(path, dpi=220); plt.close(fig); paths.append(path)
    stress = [x for x in results if x["id"] in ("S0", "S1", "S3", "S4") and x.get("el") is not None]
    fig, axes = plt.subplots(1, 2, figsize=(10, 4))
    axes[0].bar([x["id"] for x in stress], [x["r_sys"] for x in stress], color="#2c6e91")
    axes[0].set_ylabel("R_sys")
    axes[1].bar([x["id"] for x in stress], [x["el"] for x in stress], color="#a64c3c")
    axes[1].set_ylabel("Ожидаемый ущерб, млн USD/год")
    fig.tight_layout(); path = directory / "stress.png"; fig.savefig(path, dpi=220); plt.close(fig); paths.append(path)
    return paths


def build_report(path, template, results, historical, issues, quality, excel_status, figures, metadata):
    doc = Document(template)
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"; normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.first_line_indent = Cm(1.25)
    for name, size in (("Heading 1", 15), ("Heading 2", 13)):
        style = doc.styles[name]
        style.font.name = "Times New Roman"; style.font.size = Pt(size)
        style.font.bold = True; style.font.color.rgb = RGBColor(23, 54, 93)
    def p(text, style=None):
        para = doc.add_paragraph(text, style)
        para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        return para
    def formula(text):
        para = doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = para.add_run(text)
        run.font.name = "Cambria Math"
        run.font.size = Pt(12)
        return para
    def heading(text, level=1): doc.add_heading(text, level)
    def table(headers, rows):
        t = doc.add_table(rows=1, cols=len(headers)); t.style = "Table Grid"
        for j, value in enumerate(headers):
            cell = t.rows[0].cells[j]
            cell.text = str(value)
            tcpr = cell._tc.get_or_add_tcPr(); shade = OxmlElement("w:shd"); shade.set(qn("w:fill"), "17365D"); tcpr.append(shade)
            for run in cell.paragraphs[0].runs: run.font.color.rgb = RGBColor(255,255,255); run.bold = True
        for row in rows:
            cells = t.add_row().cells
            for j, val in enumerate(row): cells[j].text = str(val)
        return t
    def fmt(v): return "—" if v is None else f"{v:.4f}" if isinstance(v, (int,float)) else str(v)
    byid = {x["id"]: x for x in results}
    s0, s1 = byid["S0"], byid["S1"]
    p("МИНИСТЕРСТВО НАУКИ И ВЫСШЕГО ОБРАЗОВАНИЯ РОССИЙСКОЙ ФЕДЕРАЦИИ")
    p("РХТУ имени Д.И. Менделеева · Экономический анализ наукоёмких производств")
    heading("ОТЧЁТ ПО ПРАКТИЧЕСКОЙ РАБОТЕ № 3", 0)
    p("Расчёт надёжности и рисков цепи поставок катализаторов полимеризации. Вариант Б")
    p("Студент: ____________________    Группа: ____________________")
    p("Дата расчёта: " + metadata["timestamp_utc"][:10])
    doc.add_page_break()
    heading("1. Цель и задачи работы")
    p("Оценить годовую надёжность пятистадийной цепи, ожидаемый ущерб и хвостовой риск, проверить окупаемость резервного поставщика HF при независимости и общих шоках.")
    heading("2. Объект исследования")
    p("Цепь производства катализаторов: PGM → HF → мономер → жидкий азот → мембранный реактор. Резерв HF включается параллельно только на второй стадии; PGM остаётся единой точкой отказа.")
    heading("3. Исходные данные и контроль качества входа")
    p("Первичные параметры получены из JSON и семи CSV файлами варианта Б. Обе оси корреляционных матриц нормализованы по явной карте алиасов. Проверены диапазоны, размерности, TCO, спектр матриц, границы Фреше и 50 записей журнала.")
    table(["Показатель", "Значение"], [["Число обязательных стадий", "5"], ["Число источников HF", "2"], ["Исторические инциденты", historical["count"]], ["Блокирующие проблемы входа", sum(x.get("severity")=="BLOCKER" for x in quality)]])
    heading("4. Архитектура расчётного решения")
    p("InputDataLoader формирует типизированный SupplyChainData. DataValidator проверяет данные. Аналитический модуль и SupplyChainRiskAnalyzer считают независимые и зависимые сценарии. Отдельные модули строят Excel, журнал, графики и отчёт. Состояния каждого опыта сохраняются в CSV.")
    heading("5. Методика расчёта")
    heading("5.1. Основные показатели", 2)
    formula("R_sys = ∏R_i;    P_fail = 1 − R_sys")
    formula("R_HF = 1 − Q_aQ_b − ρ√(R_aQ_aR_bQ_b)")
    p("R_i — годовая вероятность работоспособности обязательной стадии; Q_i=1−R_i. Для HF с резервом ρ — бинарная корреляция Пирсона, допустимая в границах Фреше. Источник: лекция 5, слайды 7–10, 17–23; задание 3, с. 2–4.")
    p("Средняя длительность логнормального простоя E[T]=exp(μ+σ²/2), медиана=exp(μ). Монте-Карло моделирует один годовой сценарий и один ущерб при останове. VaR — эмпирический квантиль, CVaR — среднее худшей доли выборки с дробным весом граничного атома. Источник: лекция 5, слайды 29–34; задание 3, с. 3–4.")
    heading("5.2. Стоимостные и ресурсные показатели", 2)
    formula("EL = P_fail · (C_day · E[T] + C_scrap)")
    formula("TCO = CAPEX · r_amor + OPEX_audit + ΔC_proc")
    formula("ROI = (ΔEL − TCO) / TCO")
    p("EL измеряется в млн USD/год, C_day — млн USD/сутки, C_scrap — млн USD/инцидент. EL_med использует медиану вместо E[T] и служит отдельным фиксированным сценарием. Ставка C_day уже включает упущенную маржу. Источник: лекция 5, слайды 24–25, 29–34; задание 3, с. 2–4.")
    heading("6. Результаты расчёта")
    table(["Сценарий", "R_sys", "P_fail", "EL, млн USD/год", "ROI", "Статус"],
          [[x["id"], fmt(x.get("r_sys")), fmt(x.get("p_fail")), fmt(x.get("el")), fmt(x.get("roi")), x["status"]] for x in results])
    table(["Сценарий", "VaR95", "CVaR95", "VaR99", "CVaR99"],
          [[x["id"], fmt(x.get("var95")), fmt(x.get("cvar95")), fmt(x.get("var99")), fmt(x.get("cvar99"))] for x in results])
    p(f"Для независимого моносинга R_sys={fmt(s0.get('r_sys'))}, EL={fmt(s0.get('el'))} млн USD/год. Независимый HF резерв даёт R_sys={fmt(s1.get('r_sys'))} и EL={fmt(s1.get('el'))} млн USD/год.")
    heading("7. Harness Log")
    p(f"Excel-эталон: {excel_status}. Подробные расхождения, проблемы, тесты и допущения находятся в HARNES_LOG.xlsx. Подтверждённых записей о проблемах: {len(issues)}.")
    table(["Тип", "Свидетельство", "Статус"], [[x["type"], x["description"][:110], x.get("status","UNRESOLVED")] for x in issues[:8]] or [["INFO", "Подтверждённых ошибок не обнаружено", "PASS"]])
    heading("8. Тестирование и сценарный анализ")
    p("Автоматические тесты проверяют крайние надёжности, параллельный блок, границы Фреше, нормировку четырёх состояний, TCO, CVaR с атомом на VaR, воспроизводимость и загрузку матриц. Сценарии S0–S5 рассчитаны на одном годовом горизонте. ROI S2_full сравнивается с S0_full при той же полной матрице и seed; S4 — с S3. Результаты и статусы представлены в Scenario_Results.xlsx.")
    heading("9. Визуализация результатов")
    captions = {"chain.png":"Схема пятистадийной цепи", "scenarios.png":"Сравнение надёжности, вероятности отказа и ожидаемых потерь", "correlation.png":"Надёжность HF и ROI при изменении корреляции", "loss_distributions.png":"Распределение годового ущерба с VaR99/CVaR99", "historical.png":"Исторические простои, Q–Q и эмпирическая CDF", "stress.png":"Сопоставление базового и стрессового сценариев"}
    for i, figure in enumerate(figures, 1):
        p(f"На рисунке {i} показано: {captions.get(figure.name, figure.stem)}.")
        para = doc.add_paragraph(); para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.add_run().add_picture(str(figure), width=Cm(15.5))
        caption = doc.add_paragraph(f"Рисунок {i} — {captions.get(figure.name, figure.stem)}")
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading("10. Допущения и качество данных")
    heading("10.1. Допущения модели", 2)
    p("Один годовой сценарий состояния цепи создаёт не более одного инцидента. Корреляции CSV трактуются как целевые бинарные корреляции Пирсона и калибруются в параметры нормальной копулы. При Парето α=1,3 математическое ожидание конечно, дисперсия бесконечна. Исторические 50 событий не дают частоты годовых отказов.")
    heading("10.2. Data Quality Report", 2)
    table(["Поле", "Проблема", "Статус"], [[x["field"], x["description"][:95], x["status"]] for x in quality[:15]] or [["Вход", "Критических расхождений не найдено", "PASS"]])
    heading("11. Анализ результатов и ограничения")
    p(f"Исторический простой: среднее {historical['downtime_mean']:.2f} дня, медиана {historical['downtime_median']:.2f} дня. Шапиро для log(T): p={historical['shapiro_p']:.4f}; бутстреп KS подогнанной логнормальной модели: p={historical['ks_bootstrap_p']:.4f}. Парето: порог {historical['pareto']['threshold_days']:.1f} дня, наблюдений выше порога {historical['pareto']['tail_count']}, бутстреп KS p={historical['pareto'].get('ks_bootstrap_p',float('nan')):.4f}. AIC условной логнормальной модели {historical['pareto'].get('aic_truncated_lognormal',float('nan')):.1f}, Парето {historical['pareto'].get('aic_pareto',float('nan')):.1f}. Ни один p-value не доказывает закон распределения. Средний остаток зарегистрированного ущерба относительно C_day·T+C_scrap равен {historical['residual_mean']:.2f} млн USD.")
    if historical.get("tail_sensitivity"):
        rows = [x for x in historical["tail_sensitivity"] if x["n"] == 10000]
        log_values = [x["cvar99"] for x in rows if x["duration"] == "lognormal"]
        par_values = [x["cvar99"] for x in rows if x["duration"] == "pareto"]
        p(f"Чувствительность CVaR99 при N=10000 и трёх seed: логнормальная модель {min(log_values):.1f}–{max(log_values):.1f}, Парето {min(par_values):.1f}–{max(par_values):.1f} млн USD. Для N=2000 отдельные оценки приведены в Scenario_Results.xlsx.")
    p("Неизвестны частота повторных отказов, причинная структура общих шоков и календарная привязка надёжностей. Оценки CVaR99 основаны лишь на 100 худших из 10 000 опытов; паретовский хвост особенно нестабилен. Надёжность R и эксплуатационная готовность A имеют разные определения и не отождествляются.")
    heading("12. Итоговые выводы")
    roi = s1.get("roi")
    decision = "условно рекомендовать" if roi is not None and roi > 0 else "не рекомендовать"
    p(f"Управленческое решение: {decision} резервирование HF. Для независимых поставщиков предотвращённый ущерб {fmt(s1.get('avoided'))} млн USD/год, TCO {fmt(s1.get('tco'))} млн USD/год, ROI {fmt(roi)}. Решение зависит от подтверждения независимости и стоимости резерва.")
    for key in ("S2", "S2_full", "S3", "S4"):
        if key in byid:
            x = byid[key]
            p(f"{key}: R_sys={fmt(x.get('r_sys'))}, EL={fmt(x.get('el'))} млн USD/год, ROI={fmt(x.get('roi'))}; статус {x['status']}. Корреляция и стресс PGM ограничивают защитный эффект HF.")
    if byid.get("S2_full", {}).get("roi_ci95_low") is not None:
        x = byid["S2_full"]
        p(f"Для S2_full выборочный 95% интервал ROI: [{x['roi_ci95_low']:.3f}; {x['roi_ci95_high']:.3f}]. Он учитывает ошибку Монте-Карло при фиксированных входных параметрах, но не неопределённость самих надёжностей и корреляций.")
    p(f"Хвостовой риск S0: VaR95={fmt(s0.get('var95'))}, CVaR95={fmt(s0.get('cvar95'))}, VaR99={fmt(s0.get('var99'))}, CVaR99={fmt(s0.get('cvar99'))} млн USD; разница CVaR99−VaR99={fmt((s0.get('cvar99') or 0)-(s0.get('var99') or 0))}. Одного среднего ущерба недостаточно для оценки редких тяжёлых потерь.")
    p("Источники: задание к практическому занятию 3, с. 1–4; лекция 5 ЭА, слайды 1–35; CSV/JSON варианта Б. Формулы и определения CVaR раскрыты в ARCHITECTURE.md.")
    doc.save(path)
